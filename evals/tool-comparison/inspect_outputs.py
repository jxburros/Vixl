"""Report checkable facts about tool-comparison outputs, whichever tool made them.

    python evals/tool-comparison/inspect_outputs.py runs/T03
    python evals/tool-comparison/inspect_outputs.py runs --json > facts.json
    python evals/tool-comparison/inspect_outputs.py runs/T03 --blind blind/T03

A run is a folder holding a REPORT.md; files in its subfolders (such as badges/) belong to it, and
round2/ is a run of its own. For every file it prints what can be measured without judging taste:
pixel size, color mode, dpi, transparency, color count, animation frames and timing, seams of a
tile, ICO sizes, SVG text and embedded rasters, PDF page and print boxes, fonts, color spaces and
form fields, PPTX text, pictures and notes, and video length, frame rate and audio.

--blind copies each run's deliverables to DEST/<code>/ under a random code and writes DEST/key.csv,
so outputs can be judged without knowing which tool made them.

Needs Pillow. pypdf (PDF), python-pptx (PPTX) and ffprobe (video) are used when installed.
"""

import argparse
import csv
import json
from pathlib import Path
import random
import re
import shutil
import string
import subprocess
import sys
import xml.etree.ElementTree as ET

try:
    from PIL import Image, ImageChops, ImageSequence, ImageStat
except ImportError:  # pragma: no cover - reported at run time
    sys.exit("inspect_outputs.py needs Pillow: pip install pillow")

RASTER = {".png", ".apng", ".jpg", ".jpeg", ".webp", ".gif", ".tif", ".tiff", ".bmp"}
VIDEO = {".mp4", ".webm", ".mov", ".mkv", ".m4v"}
DELIVERABLE = RASTER | VIDEO | {".svg", ".pdf", ".pptx", ".ico"}
SKIP_DIRS = {"node_modules", ".git", "__pycache__", ".venv", "venv"}


def image_facts(path):
    with Image.open(path) as im:
        facts = {"size": f"{im.width}×{im.height}", "mode": im.mode}
        dpi = im.info.get("dpi")
        if dpi:
            facts["dpi"] = round(float(dpi[0]))
        frames = getattr(im, "n_frames", 1)
        if frames > 1:
            durations = [round(frame.info.get("duration", 0)) for frame in ImageSequence.Iterator(im)]
            facts["frames"] = frames
            facts["duration_ms"] = sum(durations)
            facts["frame_ms"] = sorted(set(durations))
            facts["loop"] = im.info.get("loop", "none")
        im.seek(0)
        rgba = im.convert("RGBA")
        alpha = rgba.getchannel("A").histogram()
        pixels = im.width * im.height
        if alpha[255] != pixels:
            facts["transparent_%"] = round(100 * alpha[0] / pixels, 1)
            facts["partial_alpha_px"] = pixels - alpha[0] - alpha[255]
        if alpha[0]:
            # Count visible colors only: fully transparent pixels may carry any RGB.
            rgba.paste((0, 0, 0, 0), mask=rgba.getchannel("A").point(lambda a: 255 if a == 0 else 0))
        colors = rgba.getcolors(maxcolors=4097)
        if colors is None:
            facts["colors"] = ">4096"
        else:
            facts["colors"] = len([c for _, c in colors if c[3]])
        if "tile" in path.stem.lower():
            facts.update(seam_facts(rgba.convert("RGB")))
    return facts


def _mean_difference(a, b):
    return sum(ImageStat.Stat(ImageChops.difference(a, b)).mean) / 3


def seam_facts(im):
    """Wrap-around difference across each edge relative to the typical difference between
    neighboring rows or columns. Close to 1 means seamless; well above 2 means a visible seam."""
    w, h = im.size
    inner_x = _mean_difference(im.crop((0, 0, w - 1, h)), im.crop((1, 0, w, h))) or 1e-6
    inner_y = _mean_difference(im.crop((0, 0, w, h - 1)), im.crop((0, 1, w, h))) or 1e-6
    edge_x = _mean_difference(im.crop((w - 1, 0, w, h)), im.crop((0, 0, 1, h)))
    edge_y = _mean_difference(im.crop((0, h - 1, w, h)), im.crop((0, 0, w, 1)))
    return {"seam_x": round(edge_x / inner_x, 2), "seam_y": round(edge_y / inner_y, 2)}


def ico_facts(path):
    with Image.open(path) as im:
        sizes = sorted(im.info.get("sizes") or im.ico.sizes())
    return {"ico_sizes": [f"{w}×{h}" for w, h in sizes]}


def svg_facts(path):
    root = ET.parse(path).getroot()
    tags = {}
    fonts = set()
    for element in root.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        tags[tag] = tags.get(tag, 0) + 1
        family = element.get("font-family")
        style = element.get("style") or ""
        match = re.search(r"font-family\s*:\s*([^;]+)", style)
        if family:
            fonts.add(family.strip())
        if match:
            fonts.add(match.group(1).strip())
    facts = {
        "width": root.get("width"),
        "height": root.get("height"),
        "viewBox": root.get("viewBox"),
        "paths": tags.get("path", 0),
        "text_elements": tags.get("text", 0),
        "embedded_images": tags.get("image", 0),
    }
    if fonts:
        facts["fonts_referenced"] = sorted(fonts)
    if "@font-face" in Path(path).read_text(encoding="utf-8", errors="replace"):
        facts["font_face"] = True
    return facts


def _inches(points):
    return f"{float(points) / 72:.2f}".rstrip("0").rstrip(".")


def _resolve(obj):
    return obj.get_object() if hasattr(obj, "get_object") else obj


def _colorspace_name(cs):
    cs = _resolve(cs)
    if isinstance(cs, list) and cs:
        name = str(cs[0])
        if name == "/ICCBased":
            components = _resolve(cs[1]).get("/N")
            return {1: "ICC-Gray", 3: "ICC-RGB", 4: "ICC-CMYK"}.get(components, "ICC")
        return name.lstrip("/")
    return str(cs).lstrip("/")


def _walk_resources(resources, seen, out):
    resources = _resolve(resources) or {}
    fonts = _resolve(resources.get("/Font")) or {}
    for font in fonts.values():
        font = _resolve(font)
        if id(font) in seen:
            continue
        seen.add(id(font))
        descriptor = font.get("/FontDescriptor")
        if descriptor is None and "/DescendantFonts" in font:
            descendant = _resolve(_resolve(font["/DescendantFonts"])[0])
            descriptor = descendant.get("/FontDescriptor")
        descriptor = _resolve(descriptor) or {}
        embedded = any(key in descriptor for key in ("/FontFile", "/FontFile2", "/FontFile3"))
        out["fonts"].append((str(font.get("/BaseFont", "?")).lstrip("/"), embedded))
    xobjects = _resolve(resources.get("/XObject")) or {}
    for xobject in xobjects.values():
        xobject = _resolve(xobject)
        if id(xobject) in seen:
            continue
        seen.add(id(xobject))
        subtype = xobject.get("/Subtype")
        if subtype == "/Image":
            cs = xobject.get("/ColorSpace")
            out["images"].append((int(xobject.get("/Width", 0)), int(xobject.get("/Height", 0)),
                                  _colorspace_name(cs) if cs is not None else "?"))
        elif subtype == "/Form":
            out["forms"].append(xobject)
            _walk_resources(xobject.get("/Resources"), seen, out)


def _operator_counts(data):
    counts = {}
    for op in re.findall(rb"(?<![A-Za-z/])(rg|RG|k|K|g|G)(?![A-Za-z])", data):
        counts[op.decode()] = counts.get(op.decode(), 0) + 1
    return counts


FIELD_KINDS = {"/Tx": "text", "/Btn": "button", "/Ch": "choice", "/Sig": "signature"}


def _field_facts(reader):
    """Terminal form fields from the AcroForm tree, with inherited kind, flags, options and length."""
    out = []

    def walk(field, parent_name, inherited):
        field = _resolve(field)
        name = str(field["/T"]) if "/T" in field else None
        full = f"{parent_name}.{name}" if parent_name and name else name or parent_name
        attrs = dict(inherited)
        for key in ("/FT", "/Ff", "/MaxLen", "/Opt"):
            if key in field:
                attrs[key] = field[key]
        kids = [_resolve(kid) for kid in field.get("/Kids", [])]
        named = [kid for kid in kids if "/T" in kid]
        if named:
            for kid in named:
                walk(kid, full, attrs)
            return
        kind = FIELD_KINDS.get(attrs.get("/FT"), "?")
        flags = int(attrs.get("/Ff", 0) or 0)
        if kind == "button":
            kind = "radio" if flags & (1 << 15) else "pushbutton" if flags & (1 << 16) else "checkbox"
        elif kind == "choice":
            kind = "dropdown" if flags & (1 << 17) else "list"
        elif kind == "text" and flags & (1 << 12):
            kind = "multiline"
        entry = {"name": full, "kind": kind}
        if flags & 2:
            entry["required"] = True
        if "/Opt" in attrs:
            entry["options"] = [str(o[-1] if isinstance(o, list) else o) for o in _resolve(attrs["/Opt"])]
        elif kind == "radio":
            states = []
            for widget in kids or [field]:
                appearance = _resolve(_resolve(widget.get("/AP") or {}).get("/N") or {})
                names = appearance.keys() if hasattr(appearance, "keys") else []
                states += [str(state).lstrip("/") for state in names if state != "/Off"]
            if states:
                entry["options"] = states
        if "/MaxLen" in attrs:
            entry["max_length"] = int(attrs["/MaxLen"])
        out.append(entry)

    acroform = _resolve(reader.trailer["/Root"].get("/AcroForm"))
    for field in (_resolve(acroform.get("/Fields")) if acroform else None) or []:
        walk(field, None, {})
    return out


def _widget_order(reader):
    """Field names in each page's annotation order (the tab order when a page sets no /Tabs)."""
    order = []
    tabs = set()
    for page in reader.pages:
        if "/Tabs" in page:
            tabs.add(str(page["/Tabs"]))
        for annotation in _resolve(page.get("/Annots")) or []:
            annotation = _resolve(annotation)
            if annotation.get("/Subtype") != "/Widget":
                continue
            node, parts = annotation, []
            while node is not None:
                if "/T" in node:
                    parts.append(str(node["/T"]))
                node = _resolve(node.get("/Parent"))
            name = ".".join(reversed(parts))
            if name and (not order or order[-1] != name):
                order.append(name)
    return order, sorted(tabs)


def pdf_facts(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        return {"note": "pip install pypdf for PDF facts"}
    reader = PdfReader(str(path))
    pages = []
    totals = {"fonts": [], "images": [], "forms": []}
    seen = set()
    ops = {}
    text_chars = 0
    for page in reader.pages:
        box = page.mediabox
        info = {"size_in": f"{_inches(box.width)}×{_inches(box.height)}"}
        boxes = [name for name in ("/TrimBox", "/BleedBox") if name in page]
        if boxes:
            info["boxes"] = [b.lstrip("/") for b in boxes]
        pages.append(info)
        _walk_resources(page.get("/Resources"), seen, totals)
        contents = page.get_contents()
        if contents is not None:
            for name, count in _operator_counts(contents.get_data()).items():
                ops[name] = ops.get(name, 0) + count
        try:
            text_chars += len((page.extract_text() or "").strip())
        except Exception:  # noqa: BLE001 - a broken text layer is itself a finding
            pages[-1]["text"] = "unreadable"
    for form in totals["forms"]:
        for name, count in _operator_counts(form.get_data()).items():
            ops[name] = ops.get(name, 0) + count
    facts = {"pages": len(reader.pages)}
    sizes = sorted({p["size_in"] for p in pages})
    facts["page_size_in"] = sizes[0] if len(sizes) == 1 else sizes
    boxes = sorted({b for p in pages for b in p.get("boxes", [])})
    if boxes:
        facts["print_boxes"] = boxes
    facts["text_chars"] = text_chars
    if totals["fonts"]:
        names = sorted({name for name, _ in totals["fonts"]})
        facts["fonts"] = names[:8] + (["…"] if len(names) > 8 else [])
        missing = sorted({name for name, embedded in totals["fonts"] if not embedded})
        if missing:
            facts["fonts_not_embedded"] = missing
    if totals["images"]:
        width_in = float(reader.pages[0].mediabox.width) / 72
        largest = max(totals["images"], key=lambda image: image[0] * image[1])
        facts["images"] = len(totals["images"])
        facts["image_colorspaces"] = sorted({cs for _, _, cs in totals["images"]})
        dpi = largest[0] / width_in
        facts["largest_image"] = f"{largest[0]}×{largest[1]} (≈{dpi:.0f} dpi if page-wide)"
    vector = {"RGB": ops.get("rg", 0) + ops.get("RG", 0), "CMYK": ops.get("k", 0) + ops.get("K", 0),
              "Gray": ops.get("g", 0) + ops.get("G", 0)}
    if any(vector.values()):
        facts["vector_color_ops"] = {k: v for k, v in vector.items() if v}
    fields = _field_facts(reader)
    if fields:
        facts["form_fields"] = fields
        order, tabs = _widget_order(reader)
        facts["widget_order"] = order
        if tabs:
            facts["tabs"] = tabs
    return facts


def pptx_facts(path):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        return {"note": "pip install python-pptx for PPTX facts"}
    deck = Presentation(str(path))
    width, height = deck.slide_width, deck.slide_height
    slides = []
    for slide in deck.slides:
        text = 0
        pictures = 0
        charts = 0
        full_picture = False
        for shape in slide.shapes:
            if shape.has_text_frame:
                text += len(shape.text_frame.text.strip())
            if getattr(shape, "has_chart", False) and shape.has_chart:
                charts += 1
            if shape.shape_type == 13:  # PICTURE
                pictures += 1
                if shape.width and shape.height and shape.width * shape.height >= 0.9 * width * height:
                    full_picture = True
        notes = 0
        if slide.has_notes_slide:
            notes = len(slide.notes_slide.notes_text_frame.text.strip())
        entry = {"text_chars": text, "pictures": pictures, "notes_chars": notes}
        if charts:
            entry["native_charts"] = charts
        if full_picture and text == 0:
            entry["picture_only"] = True
        slides.append(entry)
    return {
        "slides": len(slides),
        "slide_size_in": f"{_inches(Emu(width).pt)}×{_inches(Emu(height).pt)}",
        "picture_only_slides": sum(1 for s in slides if s.get("picture_only")),
        "slides_without_notes": sum(1 for s in slides if not s["notes_chars"]),
        "per_slide": slides,
    }


def video_facts(path):
    if not shutil.which("ffprobe"):
        return {"note": "install ffmpeg (ffprobe) for video facts"}
    result = subprocess.run(
        ["ffprobe", "-v", "error", "-print_format", "json", "-show_streams", "-show_format", str(path)],
        capture_output=True, text=True, check=False,
    )
    if result.returncode:
        return {"error": result.stderr.strip()[:200]}
    data = json.loads(result.stdout)
    facts = {"duration_s": round(float(data["format"].get("duration", 0)), 3)}
    for stream in data.get("streams", []):
        if stream.get("codec_type") == "video" and "video" not in facts:
            num, _, den = stream.get("avg_frame_rate", "0/1").partition("/")
            fps = float(num) / float(den or 1) if float(den or 1) else 0
            size = f"{stream.get('width')}×{stream.get('height')}"
            facts["video"] = f"{stream.get('codec_name')} {size} {fps:.3g} fps {stream.get('pix_fmt')}"
        elif stream.get("codec_type") == "audio" and "audio" not in facts:
            facts["audio"] = f"{stream.get('codec_name')} {stream.get('sample_rate')} Hz"
    facts.setdefault("audio", "none")
    return facts


def file_facts(path):
    suffix = path.suffix.lower()
    try:
        if suffix in RASTER:
            return image_facts(path)
        if suffix == ".ico":
            return ico_facts(path)
        if suffix == ".svg":
            return svg_facts(path)
        if suffix == ".pdf":
            return pdf_facts(path)
        if suffix == ".pptx":
            return pptx_facts(path)
        if suffix in VIDEO:
            return video_facts(path)
        if suffix == ".json":
            data = json.loads(path.read_text(encoding="utf-8"))
            return {"json": type(data).__name__, "entries": len(data) if hasattr(data, "__len__") else 1}
        if suffix == ".vixl":
            return {"kind": "Vixl document (editable)"}
    except Exception as error:  # noqa: BLE001 - an unreadable deliverable is a finding, not a crash
        return {"error": f"{type(error).__name__}: {error}"[:200]}
    return {}


def runs_under(root):
    """Map each run folder (one with a REPORT.md, else the file's own folder) to its files."""
    if root.is_file():
        return {root.parent: [root]}
    runs = {}
    files = [p for p in sorted(root.rglob("*")) if p.is_file() and not SKIP_DIRS & set(p.parts)]
    for path in files:
        run = path.parent
        for parent in path.parents:
            if (parent / "REPORT.md").is_file():
                run = parent
                break
            if parent == root:
                break
        runs.setdefault(run, []).append(path)
    return runs


def compact(value):
    if isinstance(value, dict):
        return " ".join(f"{k}={compact(v)}" for k, v in value.items())
    if isinstance(value, list):
        if value and isinstance(value[0], dict):
            return "[" + "; ".join(compact(v) for v in value) + "]"
        return ",".join(str(v) for v in value)
    return str(value)


def blind(runs, dest):
    dest.mkdir(parents=True, exist_ok=True)
    rows = []
    used = set()
    order = list(runs.items())
    random.shuffle(order)
    for run, files in order:
        deliverables = [p for p in files if p.suffix.lower() in DELIVERABLE]
        if not deliverables:
            continue
        code = None
        while code is None or code in used:
            code = "".join(random.choices(string.ascii_uppercase + string.digits, k=4))
        used.add(code)
        for path in deliverables:
            target = dest / code / path.relative_to(run)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target)
        rows.append((code, str(run)))
    with open(dest / "key.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["code", "run"])
        writer.writerows(sorted(rows))
    print(f"{len(rows)} runs copied to {dest}; the code-to-run key is {dest / 'key.csv'}")
    print("Don't open the key until judging is done.")


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("paths", nargs="+", type=Path)
    parser.add_argument("--json", action="store_true", help="print all facts as JSON")
    parser.add_argument("--blind", type=Path, metavar="DEST", help="copy deliverables under random codes")
    args = parser.parse_args()

    runs = {}
    for root in args.paths:
        if not root.exists():
            sys.exit(f"{root} does not exist")
        for run, files in runs_under(root).items():
            runs.setdefault(run, []).extend(files)

    if args.blind:
        blind(runs, args.blind)
        return

    report = {}
    for run, files in sorted(runs.items()):
        entry = {"REPORT.md": (run / "REPORT.md").is_file()}
        for path in files:
            if path.name == "REPORT.md":
                continue
            entry[str(path.relative_to(run))] = file_facts(path)
        report[str(run)] = entry

    if args.json:
        json.dump(report, sys.stdout, indent=2, ensure_ascii=False)
        print()
        return
    for run, entry in report.items():
        print(run)
        if not entry.pop("REPORT.md"):
            print("  (no REPORT.md)")
        for name, facts in entry.items():
            print(f"  {name}: {compact(facts) if facts else '-'}")
        print()


if __name__ == "__main__":
    main()
