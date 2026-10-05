"""Build Tidewick Cafe staff-meeting deck as deck.pptx (python-pptx) and deck.pdf (reportlab).

One layout spec (inches, 13.333 x 7.5) drives both renderers so the two files match.
Usage: python3 build_deck.py
"""
import os, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = 13.333, 7.5

NAVY, AMBER, FOAM, CORAL, CREAM = "14263b", "f2a541", "a8d5c8", "e2725b", "f7f1e5"
WHITE = "ffffff"
NAVY_SOFT = "4a5a6c"   # navy mixed toward cream, for secondary text on cream

# ---- shared type scale / positions for every content slide ----
MARGIN = 0.8
TITLE_BOX = (MARGIN, 0.55, W - 2 * MARGIN, 0.8)
TITLE_SIZE = 36
RULE = (MARGIN, 1.42, 1.1, 0.07)
BODY_SIZE = 28
BULLET_GAP = 40
FOOTER_SIZE = 12
FOOTER_LINE_Y = 6.82
FOOTER_Y = 6.92

# ---------------------------------------------------------------- content
SLIDES = [
    dict(kind="title",
         title="Tidewick Café — 2025 in Review",
         subtitle="Staff meeting · January 2026",
         notes="Welcome, everyone, and thank you for coming in before opening. "
               "This morning we'll look back at 2025 together: what the numbers say, what worked, "
               "and where we're heading in 2026. It should take about fifteen minutes, with time for questions at the end."),
    dict(kind="stats", title="By the numbers",
         stats=[("35,280", "cups poured", AMBER),
                ("312", "open-mic performers", FOAM),
                ("4.8★", "average review", CORAL)],
         notes="Three numbers sum up our year. We poured 35,280 cups, which is close to a hundred a day, "
               "and 312 performers took the open-mic stage. Our reviews averaged 4.8 stars, and that one belongs to every person in this room."),
    dict(kind="chart", title="Cups by quarter",
         caption="Steady all year, with a summer peak in Q3",
         data=[("Q1", 8560), ("Q2", 8830), ("Q3", 8990), ("Q4", 8900)],
         notes="Here are the cups by quarter: 8,560 in Q1, 8,830 in Q2, 8,990 in Q3 and 8,900 in Q4. "
               "The bars start at zero, so you can see how even the year really was. "
               "Q3 was our best quarter, helped by cold brew and the summer crowds on the quay."),
    dict(kind="bullets", title="What worked",
         bullets=["Thursday open mics doubled evening sales",
                  "Cold brew on tap from April to September",
                  "A pre-order pickup shelf shortened morning queues"],
         marker=FOAM,
         notes="A few things clearly worked. Thursday open mics doubled our evening sales, and cold brew on tap carried us from April through September. "
               "The pre-order pickup shelf shortened the morning queue, so thank you for keeping it tidy and labelled."),
    dict(kind="bullets", title="2026 priorities",
         bullets=["A second espresso machine for the weekend rush",
                  "A monthly makers' market on the quay",
                  "Launch the winter menu on December 1"],
         marker=AMBER,
         notes="For 2026 we have three priorities. A second espresso machine should take the pressure off the weekend rush, "
               "and a monthly makers' market will bring new faces to the quay. "
               "We'll launch the winter menu on December 1, so start thinking about drinks you'd like to see on it."),
    dict(kind="thanks", title="Thank you — see you on the quay.",
         notes="That's everything from me. Thank you for a wonderful 2025; none of these numbers happen without you. "
               "Questions are welcome now, and then let's open the doors and see you on the quay."),
]

# ---------------------------------------------------------------- element model
def rect(x, y, w, h, color, shape="rect"):
    return dict(t="rect", x=x, y=y, w=w, h=h, color=color, shape=shape)

def text(x, y, w, h, s, size, color, bold=False, align="l", anchor="t", role=None):
    return dict(t="text", x=x, y=y, w=w, h=h, s=s, size=size, color=color,
                bold=bold, align=align, anchor=anchor, role=role)

def build_elements(i, sd):
    n = i + 1
    els = []
    dark = sd["kind"] in ("title", "thanks")
    els.append(rect(0, 0, W, H, NAVY if dark else CREAM))
    if dark:
        # harbor-lamp glow + water lines
        els.append(rect(10.2, -1.3, 4.2, 4.2, AMBER, "oval"))
        els.append(rect(10.95, -0.55, 2.7, 2.7, "f6c27a", "oval"))
        for k, (yy, ww, col) in enumerate([(6.05, 5.4, FOAM), (6.3, 3.9, FOAM), (6.55, 2.4, CORAL)]):
            els.append(rect(MARGIN, yy, ww, 0.06, col))
    if sd["kind"] == "title":
        els.append(rect(MARGIN, 2.35, 1.1, 0.07, AMBER))
        els.append(text(MARGIN, 2.6, 7.0, 1.9, sd["title"], 54, CREAM, bold=True, role="title"))
        els.append(text(MARGIN, 4.55, 9.0, 0.6, sd["subtitle"], 24, FOAM, role="subtitle"))
    elif sd["kind"] == "thanks":
        els.append(rect(MARGIN, 2.55, 1.1, 0.07, AMBER))
        els.append(text(MARGIN, 2.8, 9.4, 1.9, sd["title"], 54, CREAM, bold=True, role="title"))
    else:
        els.append(text(*TITLE_BOX, sd["title"], TITLE_SIZE, NAVY, bold=True, anchor="t", role="title"))
        els.append(rect(*RULE, AMBER))
    if sd["kind"] == "stats":
        cw, gap, top, ch = 3.6, 0.367, 2.3, 3.0
        for k, (num, label, col) in enumerate(sd["stats"]):
            x = MARGIN + k * (cw + gap)
            els.append(rect(x, top, cw, ch, WHITE))
            els.append(rect(x, top, cw, 0.14, col))
            els.append(text(x + 0.35, top + 0.65, cw - 0.7, 1.2, num, 66, NAVY, bold=True))
            els.append(text(x + 0.35, top + 1.95, cw - 0.7, 0.7, label, 22, NAVY_SOFT))
    if sd["kind"] == "chart":
        els.append(text(MARGIN, 1.7, 11.0, 0.5, sd["caption"], 20, NAVY_SOFT))
        els.append(dict(t="chart", x=MARGIN, y=2.35, w=W - 2 * MARGIN, h=4.25, data=sd["data"]))
    if sd["kind"] == "bullets":
        els.append(dict(t="bullets", x=MARGIN, y=2.3, w=W - 2 * MARGIN, h=4.3,
                        items=sd["bullets"], size=BODY_SIZE, color=NAVY, marker=sd["marker"]))
    if n >= 2:  # footer on every slide after the title
        fc = FOAM if dark else NAVY_SOFT
        els.append(rect(MARGIN, FOOTER_LINE_Y, W - 2 * MARGIN, 0.015, FOAM if not dark else "2c4560"))
        els.append(text(MARGIN, FOOTER_Y, 6, 0.3, "Tidewick Café", FOOTER_SIZE, fc, bold=True, role="footer"))
        els.append(text(W - MARGIN - 2, FOOTER_Y, 2, 0.3, str(n), FOOTER_SIZE, fc, align="r", role="slidenum"))
    return els

# ---------------------------------------------------------------- PPTX renderer
FONT = "Arial"
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

def rgb(h): return RGBColor.from_string(h.upper())

def style_run(run, size, color, bold):
    f = run.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = rgb(color)

def zero_insets(tf):
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE

def pptx_text(slide, e):
    tb = slide.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    tb.name = {"title": "Title", "footer": "Footer", "slidenum": "Slide Number"}.get(e["role"], tb.name)
    tf = tb.text_frame; zero_insets(tf); tf.vertical_anchor = ANCH[e["anchor"]]
    p = tf.paragraphs[0]; p.alignment = ALIGN[e["align"]]; p.line_spacing = 1.0
    if e["role"] == "slidenum":
        # real slide-number field so it stays correct if slides are reordered
        r = p.add_run(); style_run(r, e["size"], e["color"], e["bold"]); r.text = e["s"]
        rPr = copy.deepcopy(r._r.find(qn("a:rPr")))
        fld = etree.SubElement(p._p, qn("a:fld"))
        fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"); fld.set("type", "slidenum")
        fld.append(rPr); t = etree.SubElement(fld, qn("a:t")); t.text = e["s"]
        p._p.remove(r._r)
    else:
        r = p.add_run(); r.text = e["s"]; style_run(r, e["size"], e["color"], e["bold"])
    return tb

def pptx_bullets(slide, e):
    tb = slide.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    tb.name = "Body"
    tf = tb.text_frame; zero_insets(tf)
    for k, item in enumerate(e["items"]):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.line_spacing = 1.0; p.space_after = Pt(BULLET_GAP)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.55))); pPr.set("indent", str(-Inches(0.55)))
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(buClr, qn("a:srgbClr")).set("val", e["marker"].upper())
        etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "■")
        r = p.add_run(); r.text = item; style_run(r, e["size"], e["color"], False)

def pptx_chart(slide, e):
    cd = CategoryChartData()
    cd.categories = [c for c, _ in e["data"]]
    cd.add_series("Cups poured", [v for _, v in e["data"]], number_format="#,##0")
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(e["x"]), Inches(e["y"]),
                                Inches(e["w"]), Inches(e["h"]), cd)
    gf.name = "Cups by quarter chart"
    ch = gf.chart
    ch.has_legend = False; ch.has_title = False
    ch.font.name = FONT; ch.font.size = Pt(18); ch.font.color.rgb = rgb(NAVY)
    plot = ch.plots[0]; plot.gap_width = 70; plot.vary_by_categories = False
    s = plot.series[0]
    s.format.fill.solid(); s.format.fill.fore_color.rgb = rgb(NAVY)
    best = max(range(len(e["data"])), key=lambda k: e["data"][k][1])
    pt = s.points[best]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(AMBER)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = "#,##0"; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(22); dl.font.bold = True; dl.font.color.rgb = rgb(NAVY); dl.font.name = FONT
    va = ch.value_axis
    va.minimum_scale = 0; va.maximum_scale = 10000
    va.has_major_gridlines = False; va.visible = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(20); ca.tick_labels.font.bold = True
    ca.tick_labels.font.color.rgb = rgb(NAVY_SOFT)
    ca.format.line.color.rgb = rgb(NAVY_SOFT)
    from pptx.enum.chart import XL_TICK_MARK
    ca.major_tick_mark = XL_TICK_MARK.NONE

def build_pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for i, sd in enumerate(SLIDES):
        slide = prs.slides.add_slide(blank)
        for e in build_elements(i, sd):
            if e["t"] == "rect":
                shp = slide.shapes.add_shape(MSO_SHAPE.OVAL if e["shape"] == "oval" else MSO_SHAPE.RECTANGLE,
                                             Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(e["color"]); shp.line.fill.background()
                shp.shadow.inherit = False
            elif e["t"] == "text": pptx_text(slide, e)
            elif e["t"] == "bullets": pptx_bullets(slide, e)
            elif e["t"] == "chart": pptx_chart(slide, e)
        slide.notes_slide.notes_text_frame.text = sd["notes"]
    prs.core_properties.title = "Tidewick Café — 2025 in Review"
    prs.core_properties.author = "Tidewick Café"
    prs.save(path)

# ---------------------------------------------------------------- PDF renderer
FD = "/usr/share/fonts/truetype/"
pdfmetrics.registerFont(TTFont("Sans", FD + "liberation/LiberationSans-Regular.ttf"))      # Arial-metric
pdfmetrics.registerFont(TTFont("Sans-Bold", FD + "liberation/LiberationSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("Sym", FD + "dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("Sym-Bold", FD + "dejavu/DejaVuSans-Bold.ttf"))
PT = 72.0

def hexc(c, h): c.setFillColorRGB(*(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)))

def segs(s, bold):
    """Split into runs; glyphs missing from Liberation Sans (★ ■) use DejaVu."""
    out, cur, curf = [], "", None
    for ch in s:
        f = ("Sym" if ch in "★■" else "Sans") + ("-Bold" if bold else "")
        if f != curf and cur: out.append((cur, curf)); cur = ""
        cur += ch; curf = f
    if cur: out.append((cur, curf))
    return out

def swidth(s, size, bold): return sum(pdfmetrics.stringWidth(t, f, size) for t, f in segs(s, bold))

def wrap(s, size, bold, maxw):
    lines, cur = [], ""
    for word in s.split(" "):
        trial = (cur + " " + word).strip()
        if swidth(trial, size, bold) <= maxw or not cur: cur = trial
        else: lines.append(cur); cur = word
    lines.append(cur); return lines

def draw_line(c, x, ybase, s, size, color, bold):
    hexc(c, color)
    for t, f in segs(s, bold):
        c.setFont(f, size); c.drawString(x, ybase, t); x += pdfmetrics.stringWidth(t, f, size)

LH = 1.15  # single line spacing for Arial ~ 1.15 em
ASC = 0.905

def pdf_text(c, e):
    x, w = e["x"] * PT, e["w"] * PT
    lines = wrap(e["s"], e["size"], e["bold"], w)
    lh = e["size"] * LH
    top = H * PT - e["y"] * PT
    for k, ln in enumerate(lines):
        base = top - e["size"] * ASC - k * lh - (lh - e["size"]) / 2 * 0  # top anchored
        lw = swidth(ln, e["size"], e["bold"])
        xx = x if e["align"] == "l" else (x + w - lw if e["align"] == "r" else x + (w - lw) / 2)
        draw_line(c, xx, base, ln, e["size"], e["color"], e["bold"])

def pdf_bullets(c, e):
    size, ind = e["size"], 0.55 * PT
    y = H * PT - e["y"] * PT
    for item in e["items"]:
        lines = wrap(item, size, False, e["w"] * PT - ind)
        for k, ln in enumerate(lines):
            base = y - size * ASC
            if k == 0: draw_line(c, e["x"] * PT, base, "■", size, e["marker"], False)
            draw_line(c, e["x"] * PT + ind, base, ln, size, e["color"], False)
            y -= size * LH
        y -= BULLET_GAP

def pdf_chart(c, e):
    data = e["data"]; n = len(data)
    x0, w = e["x"] * PT, e["w"] * PT
    top = H * PT - e["y"] * PT; bottom = top - e["h"] * PT
    axis_y = bottom + 0.5 * PT            # room for category labels
    plot_h = top - axis_y - 0.15 * PT
    slot = w / n; bw = slot / 1.7          # gap width 70%
    best = max(range(n), key=lambda k: data[k][1])
    for k, (cat, v) in enumerate(data):
        bx = x0 + k * slot + (slot - bw) / 2
        bh = plot_h * v / 10000
        hexc(c, AMBER if k == best else NAVY); c.rect(bx, axis_y, bw, bh, stroke=0, fill=1)
        lbl = f"{v:,}"
        draw_line(c, bx + (bw - swidth(lbl, 22, True)) / 2, axis_y + bh + 8, lbl, 22, NAVY, True)
        draw_line(c, bx + (bw - swidth(cat, 20, True)) / 2, axis_y - 26, cat, 20, NAVY_SOFT, True)
    c.setStrokeColorRGB(*(int(NAVY_SOFT[i:i + 2], 16) / 255 for i in (0, 2, 4)))
    c.setLineWidth(0.75); c.line(x0, axis_y, x0 + w, axis_y)

def build_pdf(path):
    c = rl_canvas.Canvas(path, pagesize=(W * PT, H * PT))
    c.setTitle("Tidewick Café — 2025 in Review"); c.setAuthor("Tidewick Café")
    for i, sd in enumerate(SLIDES):
        for e in build_elements(i, sd):
            if e["t"] == "rect":
                hexc(c, e["color"])
                X, Y, Wd, Ht = e["x"] * PT, H * PT - (e["y"] + e["h"]) * PT, e["w"] * PT, e["h"] * PT
                (c.ellipse(X, Y, X + Wd, Y + Ht, stroke=0, fill=1) if e["shape"] == "oval"
                 else c.rect(X, Y, Wd, Ht, stroke=0, fill=1))
            elif e["t"] == "text": pdf_text(c, e)
            elif e["t"] == "bullets": pdf_bullets(c, e)
            elif e["t"] == "chart": pdf_chart(c, e)
        c.showPage()
    c.save()

if __name__ == "__main__":
    build_pptx(os.path.join(HERE, "deck.pptx"))
    build_pdf(os.path.join(HERE, "deck.pdf"))
    print("ok")
