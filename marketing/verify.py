"""Audit the delivered marketing pack, including format-specific functionality.

Run from anywhere after a full build: python marketing/verify.py.
This is an artifact acceptance check; it does not test or mutate Vixl's engine.
"""

import hashlib
import json
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from vixl import Project
from vixl.forms import with_values
from vixl.project_folder import pack
from vixl.production import instantiate

HERE = Path(__file__).resolve().parent
OUT = HERE / "output"


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def records(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from records(child)
    elif isinstance(value, list):
        for child in value:
            yield from records(child)


def checks_and_masters():
    reports = list(OUT.rglob("*.check.json"))
    require(len(reports) >= 22, "Expected reports for the complete pack; run the full build")
    for path in reports:
        report = json.loads(path.read_text())
        require(report["passed"], f"Check did not pass: {path}")
        require(not any(issue["action"] == "fix" for issue in report["issues"]), f"Unfixed finding: {path}")
        require(all(suite["passed"] for suite in report["suites"].values()), f"Failed suite: {path}")
    masters = list(OUT.rglob("*.vixl"))
    for path in masters:
        project = Project.load(path)
        require(all(asset in project.assets for asset in project.state.get("fonts", {}).values()),
                f"Unembedded font: {path}")
        require(not any(row.get("type") == "link" or row.get("linked") for row in records(project.state)),
                f"External artwork link: {path}")
        for suite in project.state.get("suites", {}):
            require(project.check_suite(suite)["passed"], f"Reloaded suite failed: {path}: {suite}")
    return len(reports), len(masters)


def file_manifest():
    manifest = json.loads((OUT / "manifest.json").read_text())
    delivered = {path.relative_to(OUT).as_posix() for path in OUT.rglob("*")
                 if path.is_file() and path != OUT / "manifest.json"}
    require(delivered == set(manifest["files"]), "File inventory differs from manifest; rebuild")
    for name, details in manifest["files"].items():
        data = (OUT / name).read_bytes()
        require(len(data) == details["bytes"] and hashlib.sha256(data).hexdigest() == details["sha256"],
                f"Manifest hash mismatch: {name}")
    return len(delivered)


def pdfs_and_slides():
    fields = PdfReader(OUT / "print/creative-brief-letter.pdf").get_fields()
    require(set(fields) == {"project_name", "audience", "message", "action", "deliverables", "review"},
            "Creative brief fields missing")
    require(fields["review"]["/FT"] == "/Btn" and fields["message"]["/FT"] == "/Tx", "Wrong PDF field types")
    brief = Project.load(OUT / "print/creative-brief-letter.vixl")
    filled = with_values(brief, {"project_name": "Vixl community demo", "audience": "Agent developers",
                               "message": "Ideas become editable.", "action": "Try the working examples",
                               "deliverables": "Three social posts and a short motion piece", "review": "yes"})
    require(not np.array_equal(np.asarray(brief.render()), np.asarray(filled.render())), "Form fill changed no pixels")
    poster = PdfReader(OUT / "print/poster-tabloid.pdf").pages[0]
    require(abs(float(poster.trimbox.width) - 792) < 1 and abs(float(poster.trimbox.height) - 1224) < 1,
            "Tabloid trim dimensions incorrect")
    require(poster.bleedbox.width > poster.trimbox.width and poster.bleedbox.height > poster.trimbox.height,
            "Poster bleed boxes missing")
    require(b"/DeviceCMYK" in (OUT / "print/poster-tabloid.pdf").read_bytes(), "CMYK output missing")
    deck = Presentation(OUT / "deck/vixl-pitch-deck.pptx")
    require(len(deck.slides) == 12, "Expected 12 pitch slides")
    require(all(slide.has_notes_slide and len(slide.notes_slide.notes_text_frame.text) > 50 for slide in deck.slides),
            "Missing speaker notes")
    charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
    require(len(charts) == 1, "Pitch deck must contain one native chart")
    require(list(charts[0].series[0].values) == [2, 5, 3, 4], "Native chart data differs from source")
    with zipfile.ZipFile(OUT / "deck/vixl-pitch-deck.pptx") as archive:
        require(any(name.startswith("ppt/embeddings/") and name.endswith(".xlsx") for name in archive.namelist()),
                "Native chart workbook missing")
    require(len(PdfReader(OUT / "deck/vixl-pitch-deck.pdf").pages) == 12, "Deck PDF page count wrong")
    require(len(PdfReader(OUT / "carousel/carousel-1080x1350.pdf").pages) == 6, "Carousel PDF page count wrong")
    require(len(list((OUT / "carousel").glob("slide-*.png"))) == 6, "Stale or missing carousel images")


def campaign_source():
    source = OUT / "campaign/campaign-recipe.vixl"
    project = Project.load(source)
    request = json.loads((OUT / "campaign/production.json").read_text())
    report = json.loads((OUT / "campaign/rendered/production.json").read_text())
    require(report["status"] == "completed" and report["count"] == 3, "Campaign production incomplete")
    for index, row in enumerate(request["spec"]["rows"], 1):
        variant = instantiate(project, row)
        expected = np.asarray(variant.render().convert("RGBA"))
        actual = np.asarray(Image.open(OUT / f"campaign/post-{index}-1080x1080.png").convert("RGBA"))
        require(np.array_equal(expected, actual), f"Campaign variant {index} differs from recipe")
    with tempfile.TemporaryDirectory() as directory:
        restored_path = Path(directory) / "restored.vixl"
        pack(OUT / "campaign/source", restored_path)
        restored = Project.load(restored_path)
        require(np.array_equal(np.asarray(project.render()), np.asarray(restored.render())),
                "Source-folder roundtrip changed rendered pixels")
        # Prove the master renders outside the repository, without external paths.
        portable_path = Path(directory) / "portable.vixl"
        portable_path.write_bytes(source.read_bytes())
        require(np.array_equal(np.asarray(project.render()), np.asarray(Project.load(portable_path).render())),
                "Campaign master is not portable")


def motion_and_app():
    data = subprocess.run(["ffprobe", "-v", "error", "-show_streams", "-show_format", "-of", "json",
                           str(OUT / "motion/teaser.mp4")], check=True, capture_output=True, text=True)
    video = json.loads(data.stdout)
    stream = next(row for row in video["streams"] if row["codec_type"] == "video")
    require((stream["width"], stream["height"]) == (1080, 1080), "Teaser dimensions wrong")
    require(abs(float(video["format"]["duration"]) - 6.6) < 0.08, "Teaser duration wrong")
    gif = Image.open(OUT / "motion/teaser.gif")
    require(gif.size == (540, 540) and gif.is_animated, "GIF teaser invalid")
    manifest = json.loads((OUT / "app/package/manifest.json").read_text())
    require(set(manifest["states"]) == {"ready", "creating", "delivered"}, "App states missing")
    require(set(manifest["themes"]) == {"light", "dark"} and len(manifest["transitions"]) == 3,
            "App themes or transitions missing")
    for name, state in manifest["states"].items():
        master = Project.load(OUT / "app/package" / state["master"])
        for theme, variant in state["variants"].items():
            for field in ("animation", "poster", "reduced_motion"):
                require((OUT / "app/package" / variant[field]).is_file(), f"App {name}/{theme}/{field} missing")
            image = Image.open(OUT / "app/package" / variant["reduced_motion"]).convert("RGB")
            require(image.size == (640, 320), "App reduced-motion dimensions wrong")
            require(image.getpixel((0, 0)) == ((245, 247, 251) if theme == "light" else (37, 43, 57)),
                    "App theme failed to render")
            candidate = master.clone()
            candidate.state["variables"].update(variant["variables"])
            require(candidate.check(checks=["bounds", "flow", "contrast", "fonts"])["passed"],
                    f"Theme checks failed: {name}/{theme}")
    creating = Image.open(OUT / "app/package/assets/creating/dark.webp")
    require(creating.is_animated and creating.n_frames > 1, "Creating state has no animation")
    consumer = (OUT / "app/package/index.html").read_text()
    require("prefers-reduced-motion" in consumer and "window.VixlAnimation" in consumer, "App consumer incomplete")
    for path in OUT.rglob("*.motion-check.json"):
        require(json.loads(path.read_text())["passed"], f"Sampled motion checks failed: {path}")


def copy_and_links():
    copy = json.loads((OUT / "copy/marketing-copy.json").read_text())
    require(copy == json.loads((HERE / "copy.json").read_text()), "Publishing copy is stale")
    for post in copy["posts"]:
        require((OUT / post["asset"]).is_file() and post["caption"] and post["alt"], "Missing post asset/copy/alt")
    forbidden = re.compile(r"\b(?:0\.(?:19|20|21|22|23|24)\.\d+|release \d|updated|upgrade|open source)\b", re.I)
    public_text = json.dumps(copy) + (OUT / "copy/marketing-copy.txt").read_text()
    for path in OUT.rglob("*.vixl"):
        project = Project.load(path)
        public_text += " ".join(row.get("text", "") for row in records(project.state) if row.get("type") == "text")
    require(not forbidden.search(public_text), "Public copy contains a version/update/open-source claim")
    readme = (HERE / "README.md").read_text()
    for target in re.findall(r"\]\(([^)]+)\)", readme):
        if "://" not in target and not target.startswith("#"):
            require((HERE / target.split("#")[0]).exists(), f"README link missing: {target}")
    proof = (OUT / "proof.html").read_text()
    require("data:image/" in proof and "Unable to preview" not in proof, "Offline proof previews incomplete")
    for family in ("inter", "intertight", "jetbrainsmono"):
        require("SIL OPEN FONT LICENSE" in (OUT / "licenses" / f"{family}-OFL.txt").read_text(),
                f"Font licence missing: {family}")


def main():
    count = file_manifest()
    reports, masters = checks_and_masters()
    pdfs_and_slides()
    campaign_source()
    motion_and_app()
    copy_and_links()
    print(f"Verified {count} files, {masters} portable masters and {reports} passing design reports; "
          "PDF fields/print boxes, native PPTX data/notes, recipe/source roundtrip, motion, themes and copy pass.")


if __name__ == "__main__":
    main()
