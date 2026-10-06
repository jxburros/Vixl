"""Diagrams: text format, layout engines, connector routing, layers, in-place re-layout and checks."""

import io
import json
import os
import subprocess
import sys

import numpy as np
import pytest

from vixl import Project
from vixl import diagram_layout as L
from vixl.commands import compile_command
from vixl.diagrams import ICONS, parse_text
from vixl.errors import VixlError

FLOW = """
Start([Start]) -> Read[Read input] -> Valid{Valid?}
Valid -> Process[Process data]: yes
Valid -> Error[Show error]: no
Error -> Read
Process -> Save[(Database)] -> Notify[/Send mail/] -> End([Done])
Process -.-> End
"""
ORG = "CEO\n  CTO\n    Dev lead\n      Alice\n      Bob\n    QA lead\n  CFO\n    Accounting\n  COO\n    Sales\n    Support\n"


def make(text=FLOW, size=(1600, 1000), **options):
    p = Project(*size, "white")
    result = p.apply([{"type": "diagram-from-text", "name": "d", "text": text, **options}], detail="compact")
    return p, result


def geometry(p, name="d"):
    """(node boxes, edge point lists) of a diagram, in its group's coordinates."""
    layout = p.state["diagrams"][name]["layout"]
    return layout["nodes"], {k: v["points"] for k, v in layout["edges"].items()}


def lnodes(p, name="d"):
    layout = p.state["diagrams"][name]["layout"]
    return {k: L.LNode(k, w, h, *layout["shapes"].get(k, ["rect", 0.0]), x=x, y=y) for k, (x, y, w, h) in layout["nodes"].items()}


# ---------------------------------------------------------------------------------------------
# Text format


def test_text_format_chains_labels_kinds_and_attributes():
    parsed = parse_text(
        "# a comment\n"
        "direction: LR\n"
        "A[Alpha] -> B{Is it?} -> C([End])\n"
        "B -> D: yes\n"
        "D --> |maybe| E[(Store)] @color=#fde68a @icon=check\n"
        "E <-> F[/Data/]\n"
        "F -.-> G\n"
        "G -- H\n"
        "group Ops: A, C\n")
    nodes = {n["id"]: n for n in parsed["nodes"]}
    assert parsed["options"] == {"direction": "LR"}
    assert nodes["A"]["label"] == "Alpha" and nodes["B"]["kind"] == "decision" and nodes["C"]["kind"] == "terminator"
    assert nodes["E"]["kind"] == "database" and nodes["E"]["color"] == "#fde68a" and nodes["E"]["icon"] == "check"
    assert nodes["F"]["kind"] == "io" and nodes["A"]["group"] == "Ops" and nodes["Ops"]["kind"] == "group"
    edges = {(e["from"], e["to"]): e for e in parsed["edges"]}
    assert edges[("B", "D")]["label"] == "yes" and edges[("D", "E")]["label"] == "maybe"
    assert edges[("E", "F")]["arrow"] == "both" and edges[("F", "G")]["kind"] == "dashed"
    assert edges[("G", "H")]["kind"] == "line" and "arrow" not in edges[("A", "B")]


def test_brackets_only_set_the_shape_of_a_one_word_id():
    parsed = parse_text("Call foo(bar) -> Plain text (a note)\nA(Rounded) -> B")
    nodes = {n["id"]: n for n in parsed["nodes"]}
    assert "Call foo(bar)" in nodes and "Plain text (a note)" in nodes and nodes["A"]["label"] == "Rounded"


def test_text_format_hierarchy_by_indentation():
    parsed = parse_text(ORG)
    assert [n["id"] for n in parsed["nodes"]][:4] == ["CEO", "CTO", "Dev lead", "Alice"]
    pairs = [(e["from"], e["to"]) for e in parsed["edges"]]
    assert ("CEO", "CTO") in pairs and ("Dev lead", "Bob") in pairs and ("CFO", "Accounting") in pairs
    assert ("CTO", "CFO") not in pairs and all(e["kind"] == "line" for e in parsed["edges"])
    assert parsed["options"]["layout"] == "tree"
    # bullets and tabs work like spaces
    assert len(parse_text("- Root\n\t- Child\n\t- Other")["edges"]) == 2


def test_text_format_errors_name_the_line():
    for bad, needle in (("A -> ", "Line 1"), ("A -> B\n -> C", "Line 2"), ("A @weird=1", "unknown node attribute")):
        with pytest.raises(VixlError) as error:
            parse_text(bad)
        assert needle in str(error.value)


# ---------------------------------------------------------------------------------------------
# Layers


def test_a_diagram_is_ordinary_named_layers_in_one_group():
    p, result = make()
    info = result["diagram"]["d"]
    assert info["nodes"] == 8 and info["edges"] == 9 and info["layout"] == "layered" and info["scale"] == 1.0
    names = {layer["name"]: layer for layer in p.state["layers"]}
    group = names["d"]
    assert group["type"] == "group"
    children = [layer for layer in p.state["layers"] if layer.get("parent") == group["id"]]
    assert len(children) == len(p.state["layers"]) - 1 == info["layers"] - 1
    for node in ("Start", "Read", "Valid", "Process", "Error", "Save", "Notify", "End"):
        assert names[f"d/{node}"]["type"] == "shape" and names[f"d/{node}.label"]["type"] == "text"
    assert names["d/Valid"]["shape"] == "diamond" and names["d/Start"]["shape"] == "capsule"
    assert names["d/Save"]["shape"] == "path" and names["d/Valid->Process.label"]["text"] == "yes"
    edge = names["d/Read->Valid"]
    assert edge["shape"] == "path" and edge["stroke"] and edge["path"].startswith("M") and edge["path_view"] == [edge["width"], edge["height"]]
    assert names["d/Valid->Process"]["fill"] == names["d/Valid->Process"]["stroke"] != "transparent"  # arrowhead is filled
    image = p.render().convert("L")
    assert (np.asarray(image) < 250).sum() > 5000
    # the same layers feed every export
    assert b"<path" in p.export(format="SVG") and p.export(format="PDF", pdf_content="vector").startswith(b"%PDF")


def test_layers_are_ordered_lanes_then_edges_then_nodes_then_labels():
    text = "group Lane A: A, B\ngroup Lane B: C\nA -> B: go\nB -> C"
    p, _ = make(text, lanes=True)
    order = {layer["name"]: i for i, layer in enumerate(p.state["layers"])}
    lanes = max(order["d/Lane A"], order["d/Lane B"], order["d/Lane A.header"])
    edges = [i for n, i in order.items() if "->" in n and not n.endswith(".label")]
    nodes = [order[f"d/{n}"] for n in "ABC"]
    labels = [order[f"d/{n}.label"] for n in "ABC"]
    assert lanes < min(edges) and max(edges) < min(nodes) and max(nodes) < min(labels)
    assert order["d/A->B.label"] > max(labels) - 3 and order["d/Lane A.title"] > order["d/Lane A"]


def test_diagram_set_relays_out_in_place_and_keeps_ids_and_extras():
    p, _ = make("A -> B -> C")
    ids = {layer["name"]: layer["id"] for layer in p.state["layers"]}
    p.apply([{"type": "layer-style", "target": "d/A", "name": "drop-shadow", "settings": {"dx": 3, "dy": 3}}])
    result = p.apply([{"type": "diagram-set", "name": "d", "text": "C -> D: next\nA -> D", "nodes": [{"id": "A", "label": "First"}]}],
                     detail="compact")
    assert result["diagram"]["d"]["nodes"] == 4 and result["diagram"]["d"]["edges"] == 4
    after = {layer["name"]: layer["id"] for layer in p.state["layers"]}
    assert all(after[name] == ident for name, ident in ids.items()), "existing layers keep their IDs"
    assert {"d/D", "d/D.label", "d/C->D", "d/C->D.label", "d/A->D"} <= set(after) - set(ids)
    assert p.layer("d/A")["styles"] == {"drop-shadow": {"dx": 3, "dy": 3}}
    assert p.layer("d/A.label")["text"] == "First"
    p.apply([{"type": "diagram-set", "name": "d", "remove_nodes": ["D"]}])
    assert not any(layer["name"].startswith("d/D") or layer["name"] == "d/C->D" for layer in p.state["layers"])
    assert {layer["name"]: layer["id"] for layer in p.state["layers"]} == ids
    # moving the group is respected by the next layout
    p.apply([{"type": "move", "target": "d", "x": 100, "y": 50}])
    before = (p.layer("d")["x"], p.layer("d")["y"])
    p.apply([{"type": "diagram-set", "name": "d", "direction": "LR"}])
    assert p.state["diagrams"]["d"]["spec"]["direction"] == "LR"
    assert abs(p.layer("d")["x"] - 100) < abs(p.layer("d")["x"] - 700) and before[1] == 50


def test_removing_layers_and_undo_keep_the_record_consistent():
    p, _ = make("A -> B -> C")
    p.apply([{"type": "remove", "target": "d/B"}])
    spec = p.state["diagrams"]["d"]["spec"]
    assert [n["id"] for n in spec["nodes"]] == ["A", "C"] and spec["edges"] == []
    assert not any(layer["name"].startswith("d/B") or "->" in layer["name"] for layer in p.state["layers"])
    p.apply([{"type": "remove", "target": "d"}])
    assert "diagrams" not in p.state and p.state["layers"] == []
    p.undo()
    assert "d" in p.state["diagrams"] and len(p.state["layers"]) == 5
    p.apply([{"type": "diagram-set", "name": "d", "delete": True}])
    assert p.state["layers"] == [] and "diagrams" not in p.state


def test_replace_and_name_clashes():
    p, _ = make("A -> B")
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram-from-text", "name": "d", "text": "X -> Y"}])
    assert "already exists" in str(error.value)
    p.apply([{"type": "diagram-from-text", "name": "d", "text": "X -> Y", "replace": True}])
    assert [n["id"] for n in p.state["diagrams"]["d"]["spec"]["nodes"]] == ["X", "Y"]
    p.apply([{"type": "solid", "name": "taken"}])
    with pytest.raises(VixlError):
        p.apply([{"type": "diagram-from-text", "name": "taken", "text": "A -> B"}])


def test_errors_are_actionable():
    p = Project(800, 600, "white")
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram", "name": "d", "nodes": ["Alpha", "Beta"], "edges": [["Alpha", "Bta"]]}])
    assert error.value.details["suggestions"] == ["Beta"] and "unknown node" in str(error.value)
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram", "name": "d", "nodes": [{"id": "A", "kind": "hexagon"}]}])
    assert error.value.details["field"] == "nodes.0.kind" or "kind" in str(error.value)
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram", "name": "d", "nodes": [{"id": "A", "type": "decision"}]}])
    assert "Unknown field(s) 'type'" in str(error.value) and "kind" in error.value.details["allowed"]
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram", "name": "d", "nodes": [{"id": "A", "group": "Nope"}]}])
    assert "group" in str(error.value)
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram-set", "name": "ghost", "text": "A -> B"}])
    assert "Unknown diagram" in str(error.value)
    with pytest.raises(VixlError):
        p.apply([{"type": "diagram", "name": "d", "nodes": []}])
    assert p.state["layers"] == [], "a failed batch changes nothing"


def test_layer_budget_is_checked_up_front():
    from vixl.model import Limits

    p = Project(800, 600, "white", limits=Limits(max_layers=40))
    nodes = [f"N{i}" for i in range(30)]
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram", "name": "d", "nodes": nodes, "edges": [[a, b] for a, b in zip(nodes, nodes[1:])]}])
    assert error.value.code == "resource_limit" and "split it into several diagrams" in str(error.value)
    assert p.state["layers"] == []


# ---------------------------------------------------------------------------------------------
# Layout engines


GRAPH_EDGES = [("a", "b"), ("a", "c"), ("b", "d"), ("c", "d"), ("a", "e"), ("e", "f"), ("d", "f"), ("f", "g"), ("b", "g"),
               ("g", "h"), ("c", "h"), ("h", "a")]


def graph(sizes=None):
    names = "abcdefgh"
    nodes = [L.LNode(n, *(sizes or {}).get(n, (100 + 8 * i, 44))) for i, n in enumerate(names)]
    nodes[3] = L.LNode("d", 120, 76, "diamond")
    edges = [L.LEdge(f"{s}->{t}", s, t, (34, 18) if i % 3 == 0 else None) for i, (s, t) in enumerate(GRAPH_EDGES)]
    return nodes, edges


def overlaps(boxes, gap=0.0):
    items = list(boxes.items())
    for i, (a, (ax, ay, aw, ah)) in enumerate(items):
        for b, (bx, by, bw, bh) in items[i + 1:]:
            if ax < bx + bw + gap and bx < ax + aw + gap and ay < by + bh + gap and by < ay + ah + gap:
                return a, b
    return None


@pytest.mark.parametrize("algorithm", L.ALGORITHMS)
@pytest.mark.parametrize("routing", L.ROUTINGS)
def test_no_node_overlap_and_edges_start_and_end_on_borders(algorithm, routing):
    for direction in ("TB", "LR"):
        nodes, edges = graph()
        result = L.layout(nodes, edges, [], L.Options(algorithm=algorithm, routing=routing, direction=direction))
        assert overlaps(result.nodes, 10) is None, (algorithm, routing, direction)
        by_id = {n.id: n for n in nodes}
        for edge in edges:
            route = result.edges[edge.id]
            assert L.border_distance(by_id[edge.src], route.start) < 0.75, (algorithm, edge.id, "start")
            assert L.border_distance(by_id[edge.dst], route.end) < 0.75, (algorithm, edge.id, "end")
        assert result.size[0] > 0 and result.size[1] > 0
        for x, y, w, h in result.nodes.values():
            assert x >= 0 and y >= 0 and x + w <= result.size[0] + 1e-6 and y + h <= result.size[1] + 1e-6


@pytest.mark.parametrize("routing", ["orthogonal", "curved"])
def test_layered_edges_do_not_run_through_nodes(routing):
    for direction in ("TB", "LR", "BT", "RL"):
        nodes, edges = graph()
        result = L.layout(nodes, edges, [], L.Options(direction=direction, routing=routing))
        by_id = {n.id: n for n in nodes}
        for edge in edges:
            segments = result.edges[edge.id].segments
            for node in nodes:
                margin = 3.0 if node.id in (edge.src, edge.dst) else 1.0
                assert not L.route_hits_node(segments, by_id[node.id], margin), (direction, edge.id, node.id)


def test_layout_is_deterministic_across_hash_seeds():
    script = (
        "import json, sys\n"
        "sys.path.insert(0, %r)\n"
        "from vixl import Project\n"
        "p = Project(1400, 900, 'white')\n"
        "p.apply([{'type': 'diagram-from-text', 'name': 'd', 'text': %r}])\n"
        "layout = p.state['diagrams']['d']['layout']\n"
        "print(json.dumps([layout['nodes'], layout['edges'], layout['size'], [(l['name'], l['x'], l['y'], l['width'], l['height'], l.get('path')) for l in p.state['layers']]], sort_keys=True))\n"
    ) % (os.path.dirname(os.path.dirname(L.__file__)), FLOW)
    outputs = []
    for seed in ("0", "1", "12345"):
        env = {**os.environ, "PYTHONHASHSEED": seed}
        outputs.append(subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, env=env, check=True).stdout)
    assert outputs[0] == outputs[1] == outputs[2] and len(outputs[0]) > 1000


def test_cycles_self_loops_and_parallel_edges():
    p, result = make("A -> B -> C -> A\nB -> B: retry\nA -> B\nA -> B: again")
    info = result["diagram"]["d"]
    assert info["reversed_edges"] == ["C->A"]
    spec = p.state["diagrams"]["d"]["spec"]
    assert [e["id"] for e in spec["edges"]] == ["A->B", "B->C", "C->A", "B->B", "A->B#2", "A->B#3"]
    assert not any(key.startswith("_") for edge in spec["edges"] for key in edge), "styling is not stored in the spec"
    nodes, edges = geometry(p)
    loop = edges["B->B"]
    box = nodes["B"]
    assert max(x for x, _ in loop) > box[0] + box[2], "the self loop sticks out of its node"
    # the reversed edge still ends with its arrowhead on A (its target)
    assert p.layer("d/C->A")["shape"] == "path"
    assert edges["C->A"][-1][1] <= nodes["A"][1] + nodes["A"][3] + 1


def test_every_layout_and_direction_produces_a_valid_document():
    for layout in ("auto", "layered", "tree", "radial", "mindmap", "grid"):
        for direction in ("TB", "LR", "BT", "RL"):
            p, result = make(ORG, layout=layout, direction=direction, routing="curved" if layout == "radial" else "orthogonal")
            assert result["diagram"]["d"]["nodes"] == 11
            nodes, _ = geometry(p)
            assert overlaps({k: v for k, v in nodes.items()}, 6) is None, (layout, direction)


def test_tree_layout_centres_parents_and_levels_align():
    p, result = make(ORG)
    assert result["diagram"]["d"]["layout"] == "tree"
    nodes, edges = geometry(p)
    centre = lambda n: nodes[n][0] + nodes[n][2] / 2  # noqa: E731
    assert abs(centre("CEO") - (centre("CTO") + centre("COO")) / 2) < 80
    assert nodes["CTO"][1] == nodes["CFO"][1] == nodes["COO"][1] > nodes["CEO"][1]
    assert p.layer("d/CEO->CTO")["fill"] == "transparent", "hierarchy lines have no arrowheads"
    assert nodes["Alice"][1] > nodes["Dev lead"][1]


def test_swimlanes_keep_nodes_in_their_lanes():
    text = "group Customer: Order, Pay\ngroup Shop: Review, Ship\nOrder -> Pay -> Review -> Ship -> Order"
    for direction in ("TB", "LR"):
        p, _ = make(text, lanes=True, direction=direction)
        layout = p.state["diagrams"]["d"]["layout"]
        lanes, nodes = layout["groups"], layout["nodes"]
        horizontal = direction == "LR"
        axis, size = (1, 3) if horizontal else (0, 2)
        for lane, members in (("Customer", ("Order", "Pay")), ("Shop", ("Review", "Ship"))):
            lo, hi = lanes[lane][axis], lanes[lane][axis] + lanes[lane][size]
            for node in members:
                assert lo <= nodes[node][axis] and nodes[node][axis] + nodes[node][size] <= hi + 0.01, (direction, lane, node)
        assert lanes["Customer"][axis] + lanes["Customer"][size] <= lanes["Shop"][axis] + 0.01


def test_clusters_surround_their_members():
    p, _ = make("group Core: A, B\nA -> B\nB -> C", layout="layered")
    layout = p.state["diagrams"]["d"]["layout"]
    box = layout["groups"]["Core"]
    for node in "AB":
        x, y, w, h = layout["nodes"][node]
        assert box[0] <= x and box[1] <= y and x + w <= box[0] + box[2] and y + h <= box[1] + box[3]


def test_ports_route_around_nodes():
    p = Project(1000, 800, "white")
    p.apply([{"type": "diagram", "name": "d", "nodes": ["A", "B", "C"],
              "edges": [{"from": "A", "to": "B"}, {"from": "C", "to": "A", "from_port": "left", "to_port": "left"}]}])
    nodes, edges = geometry(p)
    start = edges["C->A"][0]
    assert abs(start[0] - nodes["C"][0]) < 0.01, "leaves C by its left side"
    assert abs(edges["C->A"][-1][0] - nodes["A"][0]) < 0.01, "arrives on A's left side"
    points = [tuple(pt) for pt in edges["C->A"]]
    assert not L.route_hits_node([("L", a, b) for a, b in zip(points, points[1:])], lnodes(p)["B"], 1.0)


def test_dashed_lines_arrows_and_colors():
    p = Project(800, 600, "white")
    p.apply([{"type": "diagram", "name": "d", "nodes": ["A", "B", "C", "D"],
              "edges": [{"from": "A", "to": "B", "kind": "dashed"}, {"from": "B", "to": "C", "kind": "line", "color": "#c00"},
                        {"from": "C", "to": "D", "arrow": "both"}]}])
    dashed, line, both = p.layer("d/A->B"), p.layer("d/B->C"), p.layer("d/C->D")
    assert dashed["path"].count("M") > 3, "dashes are separate runs"
    assert line["fill"] == "transparent" and line["stroke"] == "#c00" and line["path"].count("M") == 1
    assert both["path"].count("Z") == 2, "arrowheads at both ends"


def test_fit_to_canvas_shrinks_and_reports_unreadable_labels():
    nodes = [f"Task number {i}" for i in range(12)]
    edges = [[nodes[i], nodes[j]] for i, j in zip(range(11), range(1, 12))] + [[nodes[0], nodes[i]] for i in range(2, 12, 3)]
    p = Project(1200, 900, "white")
    result = p.apply([{"type": "diagram", "name": "d", "nodes": nodes, "edges": edges, "layout": "layered"}], detail="compact")
    info = result["diagram"]["d"]
    assert info["scale"] < 0.8 and info["label_size"] >= 9 and "warnings" not in info
    bounds = p.inspect()["layers"][0]["resolved_bounds"]
    assert bounds[0] >= 0 and bounds[1] >= 0 and bounds[0] + bounds[2] <= 1200 and bounds[1] + bounds[3] <= 900
    unrestricted = Project(1200, 900, "white")
    unrestricted.apply([{"type": "diagram", "name": "d", "nodes": nodes, "edges": edges, "layout": "layered", "fit": "none"}])
    assert unrestricted.state["diagrams"]["d"]["layout"]["scale"] == 1.0
    assert unrestricted.state["diagrams"]["d"]["layout"]["natural"][1] > 900
    cramped = Project(260, 180, "white")
    result = cramped.apply([{"type": "diagram", "name": "d", "nodes": nodes, "edges": edges}], detail="compact")
    assert any("hard to read" in w for w in result["diagram"]["d"].get("warnings", []))
    assert any(i["check"] == "diagram" and "scaled down" in i["message"] for i in cramped.check(checks=["diagram"])["issues"])


def test_area_margin_and_contain():
    p = Project(1000, 600, "white")
    p.apply([{"type": "diagram", "name": "d", "nodes": ["A", "B"], "edges": [["A", "B"]], "x": 500, "y": 100, "width": 400, "height": 300,
              "fit": "contain"}])
    box = p.inspect()["layers"][0]["resolved_bounds"]
    assert 500 <= box[0] and box[0] + box[2] <= 900 and 100 <= box[1] and box[1] + box[3] <= 400
    assert max(box[2] / 400, box[3] / 300) > 0.6, "contain scales a small diagram up to fill its area"


def test_theme_colors_and_icons():
    p = Project(1000, 600, "white")
    p.apply([{"type": "diagram", "name": "d", "theme": "dark", "background": "#0f172a", "nodes": [{"id": "A", "icon": "bolt"}, "B"],
              "edges": [["A", "B"]]}])
    assert p.layer("d/background")["fill"] == "#0f172a"
    assert p.layer("d/A.icon")["path"] == ICONS["bolt"] and p.layer("d/A.icon")["path_view"] == [24, 24]
    assert p.layer("d/A.label")["color"] == "#f1f5f9"
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram-set", "name": "d", "nodes": [{"id": "A", "icon": "rocket"}]}])
    assert "Unknown icon" in str(error.value)


def test_light_fill_gets_dark_text_and_dark_fill_gets_light_text():
    p = Project(1000, 600, "white")
    p.apply([{"type": "swatch", "name": "brand", "color": "#222222"},
             {"type": "diagram", "name": "d", "nodes": [{"id": "A", "color": "#fde68a"}, {"id": "B", "color": "#1e3a8a"}, {"id": "C", "color": "@brand"}],
              "edges": []}])
    from vixl.colors import contrast_ratio, parse

    fills = {"A": "#fde68a", "B": "#1e3a8a", "C": "#222222"}
    for node, fill in fills.items():
        assert contrast_ratio(parse(fill), parse(p.layer(f"d/{node}.label")["color"])) >= 4.5


# ---------------------------------------------------------------------------------------------
# Checks


def issues(p, *checks):
    return p.check(checks=list(checks or ["diagram"]))["issues"]


def test_a_fresh_diagram_passes_the_diagram_checks():
    for text in (FLOW, ORG, "A -> B -> C -> A\nA -> C"):
        p, _ = make(text)
        assert [i for i in issues(p) if i["severity"] == "error"] == []


def test_check_reports_overlapping_nodes_and_stale_edges():
    p, _ = make("A -> B -> C")
    b = p.layer("d/B")
    p.apply([{"type": "move", "target": "d/A", "x": b["x"] + 5, "y": b["y"] + 5}, {"type": "move", "target": "d/A.label", "x": b["x"] + 20, "y": b["y"] + 15}])
    found = issues(p)
    assert any("overlap" in i["message"] and "A" in i["layers"][0] for i in found if i["severity"] == "error")
    assert any("moved or resized" in i["message"] and i["severity"] == "warning" for i in found)
    p.apply([{"type": "diagram-set", "name": "d"}])
    assert [i for i in issues(p) if i["severity"] == "error"] == []


def test_check_reports_edges_through_nodes():
    p = Project(1000, 800, "white")
    p.apply([{"type": "diagram", "name": "d", "nodes": ["A", "B", "C"], "edges": [["A", "B"], ["B", "C"]]}])
    record = p.state["diagrams"]["d"]["layout"]
    nodes = record["nodes"]
    # route an edge straight through B (as a hand-edited or stale layout would)
    record["edges"]["A->B"]["points"] = [[nodes["A"][0] + nodes["A"][2] / 2, nodes["A"][1] + nodes["A"][3]],
                                         [nodes["C"][0] + nodes["C"][2] / 2, nodes["C"][1]]]
    found = issues(p)
    assert any("passes through" in i["message"] and "'B'" in i["message"] for i in found if i["severity"] == "error")


def test_check_reports_unreadable_labels():
    p, _ = make("A -> B")
    p.apply([{"type": "text-set", "target": "d/A.label", "color": "#f4f4f4"}])
    assert any("contrast" in i["message"] and i["severity"] == "error" for i in issues(p))
    p.apply([{"type": "diagram-set", "name": "d"}])
    assert not [i for i in issues(p) if "contrast" in i["message"]]
    p.apply([{"type": "text-set", "target": "d/B.label", "text": "A label that is far too long for the node it sits in", "size": 40}])
    assert any("does not fit inside" in i["message"] for i in issues(p))
    p.apply([{"type": "diagram-set", "name": "d", "size": 6}])
    p.apply([{"type": "scale", "target": "d", "value": 0.5}])
    assert any("too small to read" in i["message"] or "px tall" in i["message"] for i in issues(p))


def test_diagram_checks_are_part_of_the_default_check_set():
    p, _ = make("A -> B")
    p.apply([{"type": "move", "target": "d/B", "x": p.layer("d/A")["x"], "y": p.layer("d/A")["y"]}])
    report = p.check()
    assert any(i["check"] == "diagram" for i in report["issues"])
    assert "diagram" in report["checked"]["checks"]


# ---------------------------------------------------------------------------------------------
# Pages, exports, interfaces


def test_diagram_on_a_page_is_found_from_another_page():
    p = Project(800, 600, "white")
    p.apply([{"type": "page", "action": "add", "name": "one"}, {"type": "diagram", "name": "d", "nodes": ["A", "B"], "edges": [["A", "B"]]},
             {"type": "page", "action": "add", "name": "two"}])
    assert p.state["layers"] == []
    p.apply([{"type": "diagram-set", "name": "d", "nodes": ["C"], "edges": [["B", "C"]]}])
    assert p.inspect()["pages"][0]["layers"] > 6 and p.state["layers"] != [] and p.state["page"] == p.state["diagrams"]["d"]["page"]


def test_exports_draw_the_diagram():
    pypdf = pytest.importorskip("pypdf")
    p, _ = make("Alpha -> Beta: go\nBeta -.-> Gamma")
    png = np.asarray(p.render().convert("L"))
    assert (png < 128).sum() > 3000
    svg = p.export(format="SVG")
    assert svg.count(b"<path") >= 8
    report = {}
    pdf = pypdf.PdfReader(io.BytesIO(p.export(format="PDF", pdf_content="vector", report=report)))
    text = pdf.pages[0].extract_text()
    assert "Alpha" in text and "Gamma" in text and "go" in text and report["raster_fallbacks"] == {}
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    deck = pptx.Presentation(io.BytesIO(p.export(format="PPTX")))

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
            else:
                yield shape

    texts = [shape.text_frame.text for shape in walk(deck.slides[0].shapes) if shape.has_text_frame]
    assert "Alpha" in texts and "Beta" in texts and "go" in texts


def test_cli_commands_compile_to_operations():
    op = compile_command('diagram-from-text "A -> B\\nB -> C" --name flow --layout tree --direction LR --lanes --size 22')
    assert op == {"type": "diagram-from-text", "name": "flow", "text": "A -> B\nB -> C", "layout": "tree", "direction": "LR", "lanes": True, "size": 22.0}
    op = compile_command("diagram-set flow --remove-nodes A B --text 'C -> D'")
    assert op["remove_nodes"] == ["A", "B"] and op["type"] == "diagram-set" and op["text"] == "C -> D"
    op = compile_command("diagram flow --nodes '[\"A\", \"B\"]' --edges '[[\"A\", \"B\"]]' --no-arrows")
    assert op["nodes"] == ["A", "B"] and op["arrows"] is False


def test_schema_documents_and_rejects_fields():
    from vixl.schema import operation_schema

    variants = {v["properties"]["type"]["const"]: v for v in operation_schema()["properties"]["operations"]["items"]["oneOf"]}
    assert {"diagram", "diagram-set", "diagram-from-text"} <= set(variants)
    assert variants["diagram"]["properties"]["layout"]["enum"][0] == "auto"
    p = Project(400, 300, "white")
    with pytest.raises(VixlError) as error:
        p.apply([{"type": "diagram", "name": "d", "nodes": ["A"], "edges": [{"from": "A", "to": "A", "colour": "red"}]}])
    assert "colour" in str(error.value)
    # the aliases a model is likely to guess
    result = p.apply([{"type": "flowchart", "name": "f", "text": "A -> B"}], detail="compact")
    assert any("diagram" in note for note in result["normalized"])


def test_history_and_compact_summary_name_the_diagram():
    p, _ = make("A -> B")
    assert p.inspect()["diagrams"]["d"]["spec"]["nodes"][0]["id"] == "A"
    from vixl.changes import summarize

    assert summarize(p)["diagrams"] == ["d"]
    p.undo()
    assert "diagrams" not in p.state


def test_archive_round_trip_keeps_the_diagram_editable(tmp_path):
    p, _ = make("A -> B -> C")
    path = tmp_path / "flow.vixl"
    p.save(path)
    loaded = Project.load(path)
    ids = {layer["name"]: layer["id"] for layer in loaded.state["layers"]}
    loaded.apply([{"type": "diagram-set", "name": "d", "text": "C -> D"}])
    assert all({layer["name"]: layer["id"] for layer in loaded.state["layers"]}[n] == i for n, i in ids.items())
    assert json.dumps(loaded.state["diagrams"]["d"]["spec"])


def test_mcp_tools_apply_describe_and_check_diagrams(tmp_path):
    import asyncio

    from vixl.interfaces import mcp_server

    server = mcp_server(workspace=tmp_path)

    def call(name, arguments):
        async def run():
            result = await server.call_tool(name, arguments)
            content = result[0] if isinstance(result, tuple) else result
            return json.loads("".join(getattr(item, "text", "") for item in content))

        return asyncio.run(run())

    call("vixl_document_create", {"path": "d.vixl", "width": 900, "height": 600, "background": "white"})
    result = call("vixl_operations_apply", {"operations": [
        {"type": "diagram-from-text", "name": "flow", "text": "A -> B{Ok?}\nB -> C: yes\nB -> A: no", "direction": "LR"}]})
    assert result["diagram"]["flow"]["nodes"] == 3 and result["diagram"]["flow"]["crossings"] == 0
    report = call("vixl_check", {"checks": ["diagram"]})
    assert report["passed"] and report["checked"]["checks"] == ["diagram"]
    schema = call("vixl_operation_schema", {"types": ["diagram"]})
    assert schema["diagram"]["properties"]["nodes"]["items"]["anyOf"][1]["properties"]["kind"]["enum"][0] == "process"

    async def font_file():  # services do not take font files; typography roles are the way to set type
        try:
            await server.call_tool("vixl_operations_apply", {"operations": [{"type": "diagram", "name": "x", "nodes": ["A"], "font": "some.ttf"}]})
        except Exception as exc:  # noqa: BLE001
            return str(exc)

    assert "forbidden" in asyncio.run(font_file())
