"""Golden-file comparison for the visual regression suite.

References live in ``tests/visual/golden``. ``VIXL_UPDATE_GOLDEN=1`` rewrites them; on a mismatch the
actual output (and, for images, a diff) is written to ``tests/visual/_failures`` (override with
``VIXL_VISUAL_FAILURES``) and the message names the paths.
"""

import difflib
import io
import os
from pathlib import Path

import numpy as np
from PIL import Image

HERE = Path(__file__).parent
GOLDEN = HERE / "golden"
FAILURES = Path(os.environ.get("VIXL_VISUAL_FAILURES") or HERE / "_failures")

# A pixel "differs" when any channel moves by more than PIXEL_THRESHOLD (absorbs anti-aliasing noise);
# the comparison fails when more than MAX_PERCENT of pixels differ or any channel moves by more than MAX_DIFF.
PIXEL_THRESHOLD = 3
MAX_PERCENT = 0.05
MAX_DIFF = 48


def updating():
    return os.environ.get("VIXL_UPDATE_GOLDEN", "") not in ("", "0")


def to_rgba(image):
    return image.convert("RGBA")


def compare_images(expected, actual, pixel_threshold=PIXEL_THRESHOLD):
    """Return (max channel diff, percent of differing pixels, per-pixel diff) for two same-size images."""
    a = np.asarray(to_rgba(expected), dtype=np.int16)
    b = np.asarray(to_rgba(actual), dtype=np.int16)
    delta = np.abs(a - b).max(axis=2)
    changed = delta > pixel_threshold
    return int(delta.max()), float(changed.mean() * 100), delta


def diff_image(expected, actual, delta):
    """Side by side: reference | actual | amplified difference (white where identical)."""
    w, h = expected.size
    sheet = Image.new("RGBA", (w * 3 + 8, h), (255, 255, 255, 255))
    sheet.paste(to_rgba(expected), (0, 0))
    sheet.paste(to_rgba(actual), (w + 4, 0))
    heat = np.full((h, w, 3), 255, dtype=np.uint8)
    strength = np.clip(delta * 4, 0, 255).astype(np.uint8)
    heat[:, :, 1] = 255 - strength
    heat[:, :, 2] = 255 - strength
    sheet.paste(Image.fromarray(heat, "RGB").convert("RGBA"), (w * 2 + 8, 0))
    return sheet


def _fail_path(name):
    FAILURES.mkdir(parents=True, exist_ok=True)
    return FAILURES / name


def check_png(name, image, max_percent=MAX_PERCENT, max_diff=MAX_DIFF):
    """Compare ``image`` with ``golden/<name>.png``; raise AssertionError with paths on mismatch."""
    path = GOLDEN / f"{name}.png"
    image = to_rgba(image)
    if updating():
        GOLDEN.mkdir(parents=True, exist_ok=True)
        image.save(path, optimize=True)
        return
    assert path.exists(), f"No reference {path}. Generate it with VIXL_UPDATE_GOLDEN=1 pytest tests/visual"
    expected = Image.open(path).convert("RGBA")
    if expected.size != image.size:
        actual_path = _fail_path(f"{name}.actual.png")
        image.save(actual_path)
        raise AssertionError(f"{name}: size changed {expected.size} -> {image.size}; actual written to {actual_path}")
    peak, percent, delta = compare_images(expected, image)
    if peak > max_diff or percent > max_percent:
        actual_path = _fail_path(f"{name}.actual.png")
        diff_path = _fail_path(f"{name}.diff.png")
        image.save(actual_path)
        diff_image(expected, image, delta).save(diff_path)
        raise AssertionError(
            f"{name}: render differs from reference: {percent:.3f}% of pixels differ (limit {max_percent}%), "
            f"max channel diff {peak} (limit {max_diff}).\n  reference: {path}\n  actual:    {actual_path}\n"
            f"  diff (reference | actual | difference): {diff_path}\n"
            "If the change is intended, run VIXL_UPDATE_GOLDEN=1 pytest tests/visual and commit the new references."
        )


def check_text(name, text, suffix):
    """Compare a text snapshot with ``golden/<name>.<suffix>``."""
    path = GOLDEN / f"{name}.{suffix}"
    if updating():
        GOLDEN.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8", newline="\n")
        return
    assert path.exists(), f"No reference {path}. Generate it with VIXL_UPDATE_GOLDEN=1 pytest tests/visual"
    expected = path.read_text(encoding="utf-8")
    if expected != text:
        actual_path = _fail_path(f"{name}.actual.{suffix}")
        actual_path.write_text(text, encoding="utf-8", newline="\n")
        lines = list(difflib.unified_diff(expected.splitlines(), text.splitlines(), "reference", "actual", lineterm="", n=1))
        shown = "\n".join(line[:200] for line in lines[:40])
        raise AssertionError(
            f"{name}.{suffix}: snapshot differs from reference ({len(lines)} diff lines).\n{shown}\n"
            f"  reference: {path}\n  actual:    {actual_path}\n"
            "If the change is intended, run VIXL_UPDATE_GOLDEN=1 pytest tests/visual and commit the new references."
        )


def pdf_to_image(data, scale=1):
    import pypdfium2

    document = pypdfium2.PdfDocument(io.BytesIO(data))
    try:
        return document[0].render(scale=scale).to_pil()
    finally:
        document.close()


def pptx_snapshot(data):
    """Structural snapshot of a deck: per slide, each shape's kind, name, box (EMU), text and fill."""
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    out = [f"slide size {deck.slide_width}x{deck.slide_height}"]
    for index, slide in enumerate(deck.slides, 1):
        out.append(f"slide {index}")
        for shape in slide.shapes:
            fill = ""
            try:
                if shape.fill.type is not None and shape.fill.fore_color.type is not None:
                    fill = f" fill={shape.fill.fore_color.rgb}"
            except Exception:
                pass
            text = ""
            if shape.has_text_frame and shape.text_frame.text:
                text = " text=" + repr(shape.text_frame.text)
            out.append(f"  {shape.shape_type} {shape.name!r} at ({shape.left},{shape.top}) size ({shape.width}x{shape.height}){fill}{text}")
    return "\n".join(out) + "\n"
