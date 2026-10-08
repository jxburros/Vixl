"""Data-bound charts: layers kept in sync with a table, edits in place, exports and native PPTX charts."""

import io
import json
import zipfile
from xml.etree import ElementTree

import numpy as np
import pytest

from vixl import Project
from vixl.charts import format_number, nice_scale
from vixl.commands import compile_command
from vixl.errors import VixlError
from vixl.model import Limits
from vixl.schema import operation_schema, validate_operation

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
SALES = {"categories": MONTHS, "series": [{"name": "Cups", "values": [1200, 1100, 1250, 1300, 1380, 1400]}]}
DRINKS = {
    "categories": MONTHS,
    "series": [
        {"name": "Espresso", "values": [1200, 1100, 1250, 1300, 1380, 1400]},
        {"name": "Latte", "values": [800, 760, 820, 850, 900, 880]},
        {"name": "Cold brew", "values": [100, 110, 150, 220, 300, 380]},
    ],
}


def make(kind="bar", data=SALES, width=900, height=560, background="#ffffff", **options):
    p = Project(width, height, background)
    p.apply({"type": "chart", "name": "Sales", "kind": kind, "title": "Cups sold", **data, **options})
    return p


def children(p, name="Sales"):
    ident = p.layer(name)["id"]
    return [layer for layer in p.state["layers"] if layer.get("parent") == ident]


def part(p, key, name="Sales"):
    return next(layer for layer in children(p, name) if layer["chart_part"] == key)


def ids(p):
    return {layer["chart_part"]: layer["id"] for layer in children(p)}


def test_chart_is_a_group_of_ordinary_vector_layers():
    p = make("bar", DRINKS)
    group = p.layer("Sales")
    assert group["type"] == "group" and group["chart"]["kind"] == "bar"
    assert (group["width"], group["height"]) == (900, 560)
    kinds = {layer["type"] for layer in children(p)}
    assert kinds == {"shape", "text"}
    bars = [layer for layer in children(p) if layer["chart_part"].startswith("bar-")]
    assert len(bars) == 18 and all(layer["shape"] == "rectangle" for layer in bars)
    texts = {layer["text"] for layer in children(p) if layer["type"] == "text"}
    assert {"Cups sold", "Jan", "Jun", "Espresso", "Latte", "Cold brew", "0", "500"} <= texts
    assert p.layer("Sales/title")["text"] == "Cups sold"
    assert p.state["active_layer"] == group["id"]
    # An ordinary document: it validates, renders and survives a round trip through inspect().
    assert p.inspect()["layers"][-1]["chart"]["summary"]["totals"]["grand"] == sum(sum(s["values"]) for s in DRINKS["series"])


def test_bars_are_to_scale_from_a_zero_baseline():
    p = make("bar", DRINKS)
    scale = p.layer("Sales")["chart"]["summary"]["scale"]
    assert scale["min"] == 0 and scale["max"] >= 1400
    axis = part(p, "axis-line")
    per_unit = []
    for s, series in enumerate(DRINKS["series"]):
        for i, value in enumerate(series["values"]):
            bar = part(p, f"bar-{s}-{i}")
            assert bar["y"] + bar["height"] == pytest.approx(axis["y"] + 1, abs=1.5)  # every bar stands on the axis
            per_unit.append(bar["height"] / value)
    assert max(per_unit) / min(per_unit) < 1.03  # the same pixels per cup for every bar, 100 or 1,400


def test_stacked_bars_total_and_tile_exactly():
    p = make("stacked-bar", DRINKS, total_labels=True, value_labels=True)
    summary = p.layer("Sales")["chart"]["summary"]
    assert summary["totals"]["by_category"]["Jan"] == 2100
    assert summary["totals"]["by_series"]["Latte"] == sum(DRINKS["series"][1]["values"])
    assert p.layer("Sales/total Jan")["text"] == "2,100"
    segments = [part(p, f"bar-{s}-0") for s in range(3)]
    assert segments[0]["y"] == segments[1]["y"] + segments[1]["height"]  # no gap and no overlap
    assert segments[1]["y"] == segments[2]["y"] + segments[2]["height"]
    assert all(seg["x"] == segments[0]["x"] and seg["width"] == segments[0]["width"] for seg in segments)
    ratio = segments[0]["height"] / 1200
    assert segments[1]["height"] / 800 == pytest.approx(ratio, rel=0.03)
    assert segments[2]["height"] / 100 == pytest.approx(ratio, rel=0.15)


def test_percent_bars_fill_the_axis_and_label_values():
    data = {"categories": ["A", "B"], "series": [{"name": "x", "values": [30, 10]}, {"name": "y", "values": [10, 30]}]}
    p = make("percent-bar", data, value_labels=True)
    scale = p.layer("Sales")["chart"]["summary"]["scale"]
    assert (scale["min"], scale["max"]) == (0, 100)
    tops = [part(p, f"bar-1-{i}")["y"] for i in range(2)]
    assert tops[0] == tops[1]  # both stacks reach 100 %
    assert part(p, "tick-label-5")["text"] == "100%"


def test_horizontal_bars_put_the_first_category_on_top():
    p = make("horizontal-bar", DRINKS)
    first, last = part(p, "bar-0-0"), part(p, "bar-0-5")
    assert first["y"] < last["y"] and first["width"] < last["width"]
    assert part(p, "category-label-0")["x"] + part(p, "category-label-0")["width"] <= first["x"]


def test_line_and_area_follow_the_data():
    p = make("line", DRINKS, markers=True)
    line = part(p, "line-0-0")
    assert line["shape"] == "path" and line["fill"] == "transparent" and line["stroke_width"] >= 2
    marker = part(p, "marker-0-3")
    assert marker["shape"] == "ellipse"
    gap = {"categories": ["a", "b", "c", "d"], "series": [{"name": "s", "values": [1, None, 3, 4]}]}
    q = make("line", gap)
    assert {layer["chart_part"] for layer in children(q) if layer["chart_part"].startswith("line-")} == {"line-0-1"}
    assert not any(layer["chart_part"] == "marker-0-1" for layer in children(q))  # a gap has no point
    a = make("stacked-area", DRINKS)
    assert [layer["chart_part"] for layer in children(a) if layer["chart_part"].startswith("area-")] == ["area-0", "area-1", "area-2"]
    assert part(a, "area-2")["y"] < part(a, "area-0")["y"]  # stacked: the last series rides on top


def test_pie_and_donut_shares_and_wedges():
    data = {"categories": ["Espresso", "Latte", "Cold brew", "Tea"],
            "series": [{"name": "Cups", "values": [15900, 11960, 5260, 2510]}]}
    p = make("donut", data, center_text="{total}\ncups", legend="bottom")
    summary = p.layer("Sales")["chart"]["summary"]
    assert summary["totals"]["grand"] == 35630
    assert summary["shares"]["Espresso"] == pytest.approx(44.6247, abs=1e-3)
    assert sum(summary["shares"].values()) == pytest.approx(100, abs=1e-6)
    assert p.layer("Sales/slice label Espresso")["text"] == "44.6%"
    assert p.layer("Sales/slice label Tea")["text"] == "7.0%"
    assert p.layer("Sales/center text")["text"] == "35,630\ncups"
    wedges = [layer for layer in children(p) if layer["chart_part"].startswith("slice-") and "label" not in layer["chart_part"]]
    assert len(wedges) == 4 and all(layer["shape"] == "path" and "A" in layer["path"] for layer in wedges)
    # The first slice starts at 12 o'clock and runs clockwise: just right of the top of the ring is its color.
    image = np.asarray(p.render().convert("RGB")).astype(int)
    first, last = part(p, "slice-0"), part(p, "slice-3")
    assert tuple(image[first["y"] + 12, first["x"] + 8]) == (0, 114, 178)
    assert tuple(image[last["y"] + 12, last["x"] + last["width"] - 8]) == (204, 121, 167)  # the fourth ends at 12 o'clock
    pie = make("pie", data, value_labels="both")
    assert pie.layer("Sales/slice label Espresso")["text"] == "15,900 (44.6%)"


def test_one_slice_pie_is_a_full_circle():
    p = make("pie", {"categories": ["All"], "series": [{"name": "x", "values": [5]}]})
    wedge = part(p, "slice-0")
    assert wedge["path"].count("A") == 2
    image = np.asarray(p.render().convert("RGB"))
    cx, cy = wedge["x"] + wedge["width"] // 2, wedge["y"] + wedge["height"] // 2
    assert tuple(image[cy, cx]) != (255, 255, 255)


def test_geometry_keywords_place_the_chart():
    p = Project(1000, 800, "#fff")
    p.apply({"type": "chart", "name": "C", "x": "center", "y": "10%", "width": "50%", "height": "60%", **SALES})
    group = p.layer("C")
    assert (group["width"], group["height"]) == (500, 480)
    assert (group["x"], group["y"]) == (250, 80)


def test_edge_cases_still_draw():
    zero = make("bar", {"categories": ["a", "b"], "series": [{"name": "s", "values": [0, 0]}]})
    assert not any(layer["chart_part"].startswith("bar-") for layer in children(zero))
    assert zero.layer("Sales")["chart"]["summary"]["scale"]["max"] == 1
    gaps = make("area", {"categories": ["a", "b", "c"], "series": [{"name": "s", "values": [1, None, 3]}]}, value_labels=True)
    assert part(gaps, "area-0")["shape"] == "path" and not any(layer["chart_part"] == "value-0-1" for layer in children(gaps))
    negative = make("bar", {"categories": ["a", "b"], "series": [{"name": "s", "values": [-5, 10]}]}, value_labels=True)
    below, above = part(negative, "bar-0-0"), part(negative, "bar-0-1")
    assert below["y"] >= above["y"] + above["height"] - 2, "a negative bar hangs below the zero line"
    assert part(negative, "value-0-0")["y"] >= below["y"] + below["height"]
    one = make("line", {"categories": ["only"], "series": [{"name": "s", "values": [4]}]})
    assert part(one, "marker-0-0") and not any(layer["chart_part"].startswith("line-") for layer in children(one))
    slices = make("pie", {"categories": ["a", "b", "c"], "series": [{"name": "s", "values": [1, 0, 3]}]})
    assert part(slices, "slice-0") and part(slices, "slice-2") and not any(layer["chart_part"] == "slice-1" for layer in children(slices))
    assert slices.layer("Sales")["chart"]["summary"]["shares"]["b"] == 0


def test_dense_charts_thin_their_labels_and_skip_automatic_value_labels():
    categories = [f"Day {i + 1}" for i in range(120)]
    p = make("line", {"categories": categories, "series": [{"name": "v", "values": [i % 17 + i / 9 for i in range(120)]}]},
             width=1400, height=500)
    labels = [layer for layer in children(p) if layer["chart_part"].startswith("category-label-")]
    assert 5 < len(labels) < 40 and all("\n" not in layer["text"] for layer in labels)
    assert not any(layer["chart_part"].startswith(("value-", "marker-")) for layer in children(p))
    assert p.check(checks=["overlap"])["passed"]


CHANNELS = {"categories": MONTHS, "series": [
    {"name": "Retail", "values": [12000, 13500, 12800, 15000, 16200, 17000]},
    {"name": "Online", "values": [18000, 19500, 21000, 22800, 24000, 25100]},
    {"name": "Wholesale", "values": [6000, 6400, 7000, 6600, 7500, 8000]},
]}


def value_labels(p):
    return {layer["chart_part"]: layer for layer in children(p) if layer["chart_part"].startswith("value-")}


def test_chart_text_uses_the_line_height_table():  # #421
    from vixl.craft import LINE_HEIGHT, natural_height
    from vixl.text import font_data

    two = {"categories": MONTHS, "series": [{"name": "Espresso\nsingle origin", "values": [1, 2, 3, 4, 5, 6]},
                                            {"name": "Latte\nwith oat milk", "values": [2, 3, 4, 5, 6, 7]}]}
    p = make("bar", two, legend="right", title="Cups sold\nby month")
    parts = {layer["chart_part"]: layer for layer in children(p)}

    def pitch(layer):
        return natural_height(font_data(p, layer), layer["size"]) + layer["spacing"]

    assert pitch(parts["title"]) == pytest.approx(LINE_HEIGHT["heading"] * parts["title"]["size"], abs=1)
    first, second = parts["legend-label-0"], parts["legend-label-1"]
    assert pitch(first) == pytest.approx(LINE_HEIGHT["body"] * first["size"], abs=1)
    # A two-line legend label takes two rows: the next entry starts below it.
    assert second["y"] >= first["y"] + first["height"]


def test_automatic_value_labels_label_every_bar_or_whole_series_largest_first():  # #366
    p = make("bar", CHANNELS, legend="bottom")
    labels = value_labels(p)
    # Five-character labels are wider than a bar at the default size: all 18 get a smaller size, not just
    # the series with the shortest numbers.
    assert len(labels) == 18 and len({layer["size"] for layer in labels.values()}) == 1
    p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Jun", "series": "Online", "value": 26000}]})
    assert len(value_labels(p)) == 18
    # Two lines whose labels land on each other: the larger series keeps every label, the other gives way whole.
    close = {"categories": MONTHS, "series": [{"name": "A", "values": [100, 104, 108, 112, 116, 120]},
                                              {"name": "B", "values": [99, 103, 107, 111, 115, 119]}]}
    keys = set(value_labels(make("line", close)))
    assert keys == {f"value-0-{i}" for i in range(6)}, keys


def test_chart_internals_give_one_legibility_finding_and_no_mark_overlaps():  # #360
    data = {"categories": MONTHS, "series": [{"name": "Retail", "values": [120, 300, 90, 280, 100, 310]},
                                             {"name": "Online", "values": [200, 150, 260, 140, 270, 150]}]}
    for kind in ("line", "area", "bar"):
        p = make(kind, data, value_labels=True)
        report = p.check(checks=["overlap", "legibility"])
        assert not [x for x in report["issues"] if x["check"] == "overlap"], (kind, report["issues"])
        legibility = [x for x in report["issues"] if x["check"] == "legibility"]
        assert len(legibility) == 1 and legibility[0]["chart"] == "Sales" and len(legibility[0]["layers"]) > 10
    # On a piece seen as a thumbnail the chart's one finding is a warning.
    p.apply({"type": "canvas", "size": "instagram-post"})
    legibility = [x for x in p.check(checks=["legibility"])["issues"] if x["check"] == "legibility"]
    assert len(legibility) == 1 and legibility[0]["severity"] == "warning"


def test_small_text_on_a_large_non_thumbnail_piece_is_one_note():  # #360
    p = Project(3000, 2000, "#ffffff")
    p.apply([{"type": "text", "name": f"note {i}", "text": "fine print", "size": 24, "color": "#000000", "x": 40,
              "y": 40 + 60 * i} for i in range(6)])
    legibility = [x for x in p.check(checks=["legibility"])["issues"] if x["check"] == "legibility"]
    assert len(legibility) == 1 and legibility[0]["severity"] == "info" and len(legibility[0]["layers"]) == 6


def test_fixing_one_number_is_one_operation_with_stable_layer_ids():
    p = make("bar", SALES, value_labels=True)
    before = ids(p)
    old = part(p, "bar-0-5")["height"]
    result = p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Jun", "value": 700}]}, detail="compact")
    assert ids(p) == before, "no layer was created, removed or re-identified"
    assert part(p, "bar-0-5")["height"] < old
    assert part(p, "value-0-5")["text"] == "700"
    assert p.layer("Sales")["chart"]["series"][0]["values"][5] == 700
    assert p.layer("Sales")["chart"]["summary"]["totals"]["by_series"]["Cups"] == 1200 + 1100 + 1250 + 1300 + 1380 + 700
    assert result["success"] and any(change.get("chart") for change in result["changes"]["layers"].values())
    # The scale follows the data: a bigger number rescales every bar, still to the same zero baseline.
    heights = {k: part(p, k)["height"] for k in ("bar-0-0", "bar-0-4")}
    p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Mar", "value": 5000}]})
    assert before.items() <= ids(p).items(), "the scale grew: new gridlines appear, nothing that was there changes identity"
    assert part(p, "bar-0-0")["height"] < heights["bar-0-0"]
    assert p.layer("Sales")["chart"]["summary"]["scale"]["max"] >= 5000


def test_edits_keep_what_was_added_to_the_generated_layers():
    p = make("bar", SALES)
    p.apply([{"type": "layer-style", "target": "Sales/bar Cups Jan", "name": "drop-shadow", "settings": {"dx": 2, "dy": 2}},
             {"type": "opacity", "target": "Sales/bar Cups Feb", "value": 0.5}])
    p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Jan", "value": 10}]})
    assert "drop-shadow" in part(p, "bar-0-0")["styles"]
    assert part(p, "bar-0-1")["opacity"] == 0.5


def test_rows_columns_and_series_can_be_added_and_removed():
    p = make("bar", DRINKS)
    before = ids(p)
    p.apply({"type": "chart-data", "target": "Sales", "append": [{"category": "Jul", "values": [1420, 870, 410]}]})
    assert len(p.layer("Sales")["chart"]["categories"]) == 7
    assert all(ids(p)[key] == value for key, value in before.items() if key.startswith("bar-"))  # old bars untouched
    assert "bar-2-6" in ids(p)
    p.apply({"type": "chart-data", "target": "Sales", "append": [{"category": "Aug", "values": {"Latte": 870, "Espresso": 1420}}]})
    assert p.layer("Sales")["chart"]["series"][2]["values"][7] is None
    p.apply({"type": "chart-data", "target": "Sales", "remove_categories": ["Jan"], "remove_series": ["Cold brew"]})
    chart = p.layer("Sales")["chart"]
    assert chart["categories"][0] == "Feb" and [s["name"] for s in chart["series"]] == ["Espresso", "Latte"]
    assert not any(layer["chart_part"].startswith("bar-2-") for layer in children(p))
    p.apply({"type": "chart-data", "target": "Sales", "add_series": [{"name": "Tea", "values": [1] * 7}]})
    assert p.layer("Sales")["chart"]["series"][2]["name"] == "Tea"
    p.apply({"type": "chart-data", "categories": ["x", "y"], "series": [{"name": "only", "values": [3, 4]}]})
    assert [s["name"] for s in p.layer("Sales")["chart"]["series"]] == ["only"]  # the active layer is the chart


def test_restyling_resizing_and_switching_the_kind_redraws_in_place():
    p = make("bar", DRINKS)
    ident = p.layer("Sales")["id"]
    p.apply({"type": "chart", "target": "Sales", "kind": "line", "title": "Trend", "width": 700, "height": 400, "x": 20, "y": 30})
    group = p.layer("Sales")
    assert group["id"] == ident and group["chart"]["kind"] == "line" and (group["width"], group["height"]) == (700, 400)
    assert (group["x"], group["y"]) == (20, 30) and p.layer("Sales/title")["text"] == "Trend"
    assert not any(layer["chart_part"].startswith("bar-") for layer in children(p))
    p.apply({"type": "chart", "target": "Sales", "title": None, "legend": "none", "gridlines": False})
    assert not any(layer["chart_part"] in ("title", "legend-label-0") or layer["chart_part"].startswith("gridline")
                   for layer in children(p))
    p.apply({"type": "rename", "target": "Sales", "name": "Revenue"})
    p.apply({"type": "chart-data", "target": "Revenue", "set": [{"category": "Jan", "series": "Latte", "value": 5}]})
    assert all(layer["name"].startswith("Revenue/") for layer in children(p, "Revenue"))


def test_duplicate_is_an_independent_chart():
    p = make("bar", SALES)
    p.apply({"type": "duplicate", "target": "Sales", "name": "Copy"})
    p.apply({"type": "chart-data", "target": "Copy", "set": [{"category": "Jan", "value": 99}]})
    assert p.layer("Sales")["chart"]["series"][0]["values"][0] == 1200
    assert p.layer("Copy")["chart"]["series"][0]["values"][0] == 99
    assert len({layer["id"] for layer in p.state["layers"]}) == len(p.state["layers"])


def test_history_save_and_load(tmp_path):
    p = make("bar", SALES)
    p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Jan", "value": 1}]})
    p.save(tmp_path / "c.vixl")
    loaded = Project.load(tmp_path / "c.vixl")
    assert loaded.layer("Sales")["chart"]["series"][0]["values"][0] == 1
    loaded.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Feb", "value": 2}]})
    loaded.undo()
    assert loaded.layer("Sales")["chart"]["series"][0]["values"][1] == 1100
    assert loaded.render().tobytes() == p.render().tobytes()


def test_rasterizing_a_chart_leaves_a_valid_document():
    p = make("bar", SALES)
    p.apply({"type": "rasterize", "target": "Sales"})
    assert p.layer("Sales")["type"] == "raster"
    with pytest.raises(VixlError, match="not a chart"):
        p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Jan", "value": 1}]})


def test_csv_binds_a_chart_and_reload_follows_the_file(tmp_path):
    (tmp_path / "data").mkdir()
    csv = tmp_path / "data" / "cups.csv"
    csv.write_text("Month,Espresso,Latte\nJan,\"1,200\",800\nFeb,1100,\nMar,1250,820\n")
    p = Project(900, 560, "#fff")
    p._workspace = tmp_path
    p.apply({"type": "chart", "name": "Sales", "kind": "bar", "csv": "data/cups.csv", "title": "From CSV"})
    chart = p.layer("Sales")["chart"]
    assert chart["categories"] == ["Jan", "Feb", "Mar"] and chart["source"] == {"csv": "data/cups.csv"}
    assert chart["series"][0]["values"] == [1200, 1100, 1250] and chart["series"][1]["values"] == [800, None, 820]
    before = ids(p)
    csv.write_text("Month,Espresso,Latte\nJan,1200,800\nFeb,1100,900\nMar,3250,820\n")
    p.apply({"type": "chart-data", "target": "Sales", "reload": True})
    assert p.layer("Sales")["chart"]["series"][0]["values"][2] == 3250 and p.layer("Sales")["chart"]["series"][1]["values"][1] == 900
    assert set(before) <= set(ids(p)) and all(ids(p)[k] == v for k, v in before.items())
    p.apply({"type": "chart", "target": "Sales", "csv": "data/cups.csv", "series_columns": ["Latte"]})
    assert [s["name"] for s in p.layer("Sales")["chart"]["series"]] == ["Latte"]
    with pytest.raises(VixlError, match="not bound"):
        make("bar", SALES).apply({"type": "chart-data", "target": "Sales", "reload": True})


def test_csv_stays_inside_the_workspace(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    (tmp_path / "secret.csv").write_text("a,b\n1,2\n")
    p = Project(300, 200)
    p._workspace = workspace
    for path in ("../secret.csv", str(tmp_path / "secret.csv")):
        with pytest.raises(VixlError) as caught:
            p.apply({"type": "chart", "csv": path})
        assert caught.value.code == "forbidden"
    with pytest.raises(VixlError, match="not found"):
        p.apply({"type": "chart", "csv": "missing.csv"})
    (workspace / "bad.csv").write_text("a,b\nx,not-a-number\n")
    with pytest.raises(VixlError, match="not a number"):
        p.apply({"type": "chart", "csv": "bad.csv"})


def test_service_sessions_resolve_csv_in_their_workspace(tmp_path):
    from vixl.interfaces import Session

    (tmp_path / "sales.csv").write_text("Quarter,Cups\nQ1,8560\nQ2,8830\nQ3,8990\nQ4,8900\n")
    session = Session(workspace=tmp_path)
    session.create("deck.vixl", 960, 540)
    session.apply([{"type": "chart", "name": "Cups", "csv": "sales.csv", "x": 40, "y": 40, "width": 600, "height": 400}])
    assert session.inspect("Cups")["chart"]["series"][0]["values"] == [8560, 8830, 8990, 8900]
    with pytest.raises(VixlError):
        session.apply([{"type": "chart", "csv": "../outside.csv"}])


def test_rest_service_draws_and_edits_charts(tmp_path):
    from fastapi.testclient import TestClient

    from vixl.interfaces import create_app

    Project(600, 400, "#ffffff").save(tmp_path / "doc.vixl")
    (tmp_path / "cups.csv").write_text("Month,Cups\nJan,10\nFeb,20\n")
    client = TestClient(create_app(tmp_path / "doc.vixl"))
    drawn = client.post("/operations", json={"operations": [{"type": "chart", "name": "S", "csv": "cups.csv"}]})
    assert drawn.status_code == 200, drawn.text
    fixed = client.post("/operations", json={"operations": [{"type": "chart-data", "target": "S", "set": [{"category": "Feb", "value": 25}]}]})
    assert fixed.status_code == 200 and Project.load(tmp_path / "doc.vixl").layer("S")["chart"]["series"][0]["values"] == [10, 25]
    assert client.get("/render").content.startswith(b"\x89PNG")
    outside = client.post("/operations", json={"operations": [{"type": "chart", "csv": "../cups.csv"}]})
    assert outside.status_code == 403


def test_table_rows_and_chartjs_spellings_are_accepted():
    p = Project(600, 400)
    p.apply({"type": "chart", "name": "T", "table": [["", "A", "B"], ["x", 1, 2], ["y", 3, 4]]})
    assert p.layer("T")["chart"]["series"][1] == {"name": "B", "values": [2, 4]}
    result = p.apply({"type": "add-chart", "name": "J", "chartType": "column", "labels": ["a", "b"],
                      "datasets": [{"label": "S", "data": [1, 2], "backgroundColor": "#ff0000"}]})
    chart = p.layer("J")["chart"]
    assert chart["kind"] == "bar" and chart["series"] == [{"name": "S", "values": [1, 2], "color": "#ff0000"}]
    assert any("labels" in note for note in result["normalized"])
    image = np.asarray(p.render().convert("RGB"))
    assert (image == (255, 0, 0)).all(axis=2).any()


def test_colors_fonts_and_text_follow_the_document():
    p = Project(900, 560, "#0f172a")
    p.apply([{"type": "palette-apply", "name": "neon"}])
    p.apply({"type": "chart", "name": "C", "kind": "bar", **DRINKS})
    colors = p.layer("C")["chart"]["summary"]["colors"]
    palette = p.state["palettes"]["neon"]
    assert set(colors[:2]) <= set(palette)
    assert part(p, "bar-0-0", "C")["fill"] == colors[0]
    ink = part(p, "tick-label-0", "C")["color"]
    assert ink != "#1f2937", "text turns light on a dark background"
    q = Project(900, 560, "#ffffff")
    q.apply([{"type": "swatch", "name": "ink", "color": "#102030"}, {"type": "swatch", "name": "accent", "color": "#ff6600"}])
    q.apply({"type": "chart", "name": "C", "kind": "bar", "title": "T", **SALES})
    assert part(q, "bar-0-0", "C")["fill"] == "@accent" and part(q, "title", "C")["color"] == "@ink"
    q.apply({"type": "swatch", "name": "accent", "color": "#00aa00"})
    image = np.asarray(q.render().convert("RGB"))
    assert (image == (0, 170, 0)).all(axis=2).any(), "a swatch change recolors the chart"
    assert part(q, "title", "C")["font_role"] == "heading" and part(q, "tick-label-0", "C")["font_role"] == "body"


def test_text_sizes_sit_on_the_document_type_scale():
    p = Project(900, 560, "#fff")
    for name, size in (("small", 14), ("body", 18), ("lead", 24), ("h2", 32), ("h1", 48)):
        p.apply({"type": "style-define", "name": name, "settings": {"size": size}})
    p.apply({"type": "chart", "name": "C", "kind": "donut", "title": "T", "subtitle": "s", "center_text": "{total} cups",
             "categories": ["a", "b"], "series": [{"name": "s", "values": [3, 4]}], "value_labels": True})
    sizes = {layer["size"] for layer in children(p, "C") if layer["type"] == "text"}
    assert sizes <= {14, 18, 24, 32, 48}, sizes
    q = Project(900, 560, "#fff")
    q.apply({"type": "chart", "name": "C", "kind": "bar", "font_size": 17, **SALES})
    assert part(q, "tick-label-0", "C")["size"] == 17


def test_number_format_scale_and_labels_stay_in_sync():
    p = make("bar", SALES, number_format='#,##0,"K"', value_labels=True, ticks=3, max=2000)
    assert part(p, "tick-label-0")["text"] == "0K" or part(p, "tick-label-0")["text"] == "0"
    scale = p.layer("Sales")["chart"]["summary"]["scale"]
    assert scale["max"] == 2000
    labels = [layer["text"] for layer in children(p) if layer["chart_part"].startswith("value-")]
    assert labels and all(text.endswith("K") for text in labels)
    q = make("bar", {"categories": ["a", "b"], "series": [{"name": "s", "values": [0.25, 0.5]}]}, number_format="0%")
    assert part(q, "value-0-1")["text"] == "50%" and part(q, "tick-label-1")["text"].endswith("%")


def test_format_numbers_like_a_spreadsheet():
    assert format_number(1234567.891, "#,##0") == "1,234,568"
    assert format_number(0.125, "0.00") == "0.13"
    assert format_number(-1234.5, "$#,##0.0") == "-$1,234.5"
    assert format_number(0.4462, "0.0%") == "44.6%"
    assert format_number(1500, '#,##0,"K"') == "2K" and format_number(1500, "0.0,") == "1.5"
    assert format_number(7, "000") == "007" and format_number(-0.001, "0.0") == "0.0"
    with pytest.raises(VixlError, match="Unsupported number_format"):
        format_number(1, "yyyy-mm-dd")


def test_nice_scale_picks_round_steps():
    assert nice_scale(0, 1500, 5)[:3] == (0, 1500, 500)
    assert nice_scale(0, 12, 5)[:3] == (0, 12, 2)
    assert nice_scale(-90, 150, 5)[2] == 50
    assert nice_scale(0, 0, 5)[:2] == (0, 1)
    assert nice_scale(0, 3330, 5)[1] == 3500
    low, high, step, ticks = nice_scale(0, 10, 5, fixed_min=0, fixed_max=10)
    assert (low, high, ticks[-1]) == (0, 10, 10)


@pytest.mark.parametrize("op, message", [
    ({"type": "chart"}, "needs data"),
    ({"type": "chart", "kind": "bar", "categories": ["a", "a"], "series": [{"name": "s", "values": [1, 2]}]}, "unique"),
    ({"type": "chart", "categories": ["a", "b"], "series": [{"name": "s", "values": [1]}]}, "one per category"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": ["x"]}]}, "must be"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}, {"name": "s", "values": [2]}]}, "unique"),
    ({"type": "chart", "kind": "pie", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}, {"name": "t", "values": [1, 2]}]}, "one series"),
    ({"type": "chart", "kind": "pie", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, -2]}]}, "negative"),
    ({"type": "chart", "kind": "pie", "categories": ["a"], "series": [{"name": "s", "values": [0]}]}, "positive total"),
    ({"type": "chart", "kind": "percent-bar", "categories": ["a"], "series": [{"name": "s", "values": [-1]}]}, "negative"),
    ({"type": "chart", "kind": "area", "categories": ["a"], "series": [{"name": "s", "values": [1]}]}, "two categories"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "number_format": "hh:mm"}, "number_format"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "min": 5, "max": 1}, "less than max"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "title_font": "Comic Nope"}, "title_font"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "value_labels": "percent"}, "pie and donut"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "center_text": "x"}, "donut"),
    ({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "colors": ["notacolor"]}, "color"),
    ({"type": "chart", "categories": ["a"], "table": [["", "s"], ["a", 1]]}, "once"),
    ({"type": "chart-data", "target": "Sales", "set": [{"category": "Decmber", "value": 1}]}, "Unknown category"),
    ({"type": "chart-data", "target": "Sales", "set": [{"category": "Jan", "series": "Nope", "value": 1}]}, "Unknown series"),
    ({"type": "chart-data", "target": "Sales"}, "something to change"),
])
def test_bad_charts_are_rejected_with_a_reason(op, message):
    p = make("bar", SALES)
    before = p.render().tobytes()
    with pytest.raises(VixlError, match=message):
        p.apply(op)
    assert p.render().tobytes() == before, "a rejected chart operation changes nothing"


def test_a_tampered_recipe_is_refused():
    p = make("bar", SALES)
    p.layer("Sales")["chart"]["summary"]["legend"] = "diagonal"
    with pytest.raises(VixlError, match="Invalid chart summary"):
        p.apply({"type": "hide", "target": "Sales"})
    q = make("bar", SALES)
    q.layer("Sales")["chart"]["kind"] = "radar"
    with pytest.raises(VixlError, match="kind"):
        q.apply({"type": "hide", "target": "Sales"})


def test_unknown_category_suggests_the_closest():
    p = make("bar", {"categories": ["November", "December"], "series": [{"name": "s", "values": [1, 2]}]})
    with pytest.raises(VixlError) as caught:
        p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Decmber", "value": 1}]})
    assert "December" in caught.value.details["suggestions"]


def test_a_chart_too_big_for_the_layer_limit_is_refused():
    categories = [f"c{i}" for i in range(200)]
    p = Project(2000, 600, limits=Limits(max_layers=512))
    with pytest.raises(VixlError) as caught:
        p.apply({"type": "chart", "categories": categories, "series": [{"name": s, "values": [1] * 200} for s in "abc"]})
    assert caught.value.code == "resource_limit"


def test_chart_is_too_small_for_its_labels():
    with pytest.raises(VixlError, match="too small"):
        make("bar", DRINKS, width=60, height=60)


def test_checks_see_inside_the_chart():
    p = make("stacked-bar", DRINKS, value_labels=True, font_size=26)
    report = p.check(checks=["overlap", "contrast", "bounds"])
    assert report["passed"], report["issues"]
    assert report["checked"]["text_layers"] > 20
    low = make("bar", SALES, text_color="#d8d8d8")
    issues = low.check(checks=["contrast"])["issues"]
    assert issues and all(issue["layers"][0].startswith("Sales/") for issue in issues)
    legibility = make("bar", SALES, font_size=8).check(checks=["legibility"], thumbnail_width=900)
    assert any(issue["layers"][0].startswith("Sales/") for issue in legibility["issues"])


def test_chart_contrast_matches_the_general_measurement():
    from vixl.charts import chart_contrast
    from vixl.checks import canvas_projection
    from vixl.render import resolve_layout, resolved_layers

    p = make("stacked-bar", DRINKS, value_labels=True)
    resolved = {layer["id"]: layer for layer in resolved_layers(p)}
    local = resolve_layout(p, layers=list(resolved.values()))
    bounds = canvas_projection(resolved, local)["bounds"]
    texts = [layer for layer in resolved.values() if layer["type"] == "text"]
    fast = chart_contrast(p, resolved, local, bounds, texts)
    assert len(fast) == len(texts)
    for item in (texts[0], texts[len(texts) // 2], texts[-1]):
        slow = p.measure(target=item["id"], histogram="none")["contrast"]
        assert fast[item["id"]]["p10"] == pytest.approx(slow["p10"], abs=0.01)
        assert fast[item["id"]]["minimum"] == pytest.approx(slow["minimum"], abs=0.01)


def test_png_render_draws_the_data():
    p = make("bar", {"categories": ["a", "b"], "series": [{"name": "s", "values": [10, 20]}]}, colors=["#ff0000"], legend="none",
             gridlines=False, title=None, value_labels=False)
    image = np.asarray(p.render().convert("RGB")).astype(int)
    red = (image == (255, 0, 0)).all(axis=2)
    columns = np.flatnonzero(red.any(axis=0))
    left, right = columns[columns < 450], columns[columns >= 450]
    heights = [red[:, c].sum() for c in (left[len(left) // 2], right[len(right) // 2])]
    assert heights[1] / heights[0] == pytest.approx(2, rel=0.03)


def test_vector_exports_keep_the_chart_as_paths_and_text(tmp_path):
    p = make("donut", {"categories": ["a", "b", "c"], "series": [{"name": "s", "values": [1, 2, 3]}]})
    svg = p.export(tmp_path / "c.svg", format="SVG")
    assert svg.count(b"<path") >= 3 + 1 and b"<image" not in svg
    report = {}
    pdf = p.export(tmp_path / "c.pdf", format="PDF", pdf_content="vector", report=report)
    assert pdf.startswith(b"%PDF") and report["raster_fallbacks"] == {}
    pypdf = pytest.importorskip("pypdf")
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    assert "Cups sold" in reader.pages[0].extract_text()
    png = p.export(tmp_path / "c.png", format="PNG")
    assert png.startswith(b"\x89PNG")


def deck(kind="bar", data=DRINKS, **options):
    p = Project(1280, 720, "#ffffff")
    p.apply([{"type": "page", "action": "add", "name": "results"},
             {"type": "text", "text": "Results", "name": "title", "size": 44, "x": 60, "y": 30, "color": "#111"},
             {"type": "chart", "name": "Sales", "kind": kind, "x": 60, "y": 110, "width": 1160, "height": 560,
              "title": "Cups per month", "number_format": "#,##0", **data, **options}])
    return p


def chart_shapes(data):
    pptx = pytest.importorskip("pptx")
    presentation = pptx.Presentation(io.BytesIO(data))
    return [shape for slide in presentation.slides for shape in slide.shapes if shape.has_chart]


@pytest.mark.parametrize("kind, chart_type", [
    ("bar", "COLUMN_CLUSTERED"), ("stacked-bar", "COLUMN_STACKED"), ("percent-bar", "COLUMN_STACKED_100"),
    ("horizontal-bar", "BAR_CLUSTERED"), ("stacked-horizontal-bar", "BAR_STACKED"), ("line", "LINE_MARKERS"),
    ("area", "AREA"), ("stacked-area", "AREA_STACKED"),
])
def test_pptx_export_has_a_native_chart_with_the_data(kind, chart_type):
    p = deck(kind)
    report = {}
    data = p.export(format="PPTX", report=report)
    assert data == p.export(format="PPTX"), "PowerPoint output is deterministic"
    (shape,) = chart_shapes(data)
    chart = shape.chart
    assert chart.chart_type.name == chart_type
    plot = chart.plots[0]
    assert list(plot.categories) == MONTHS
    assert [s.name for s in plot.series] == ["Espresso", "Latte", "Cold brew"]
    assert [list(s.values) for s in plot.series] == [list(map(float, s["values"])) for s in DRINKS["series"]]
    assert chart.has_title and chart.has_legend and shape.name == "Sales"
    assert report["charts"] == {"1": [{"layer": "Sales", "native": True}]} and report["raster_fallbacks"] == {}
    assert chart.value_axis.minimum_scale == 0 or kind.startswith("percent")
    assert (shape.left, shape.top) == (round(60 * 6858000 / 720), round(110 * 6858000 / 720))


def test_pptx_pie_and_donut_charts_carry_slices_colors_and_labels():
    data = {"categories": ["Espresso", "Latte", "Tea"], "series": [{"name": "Cups", "values": [15900, 11960, 2510]}]}
    for kind, chart_type in (("pie", "PIE"), ("donut", "DOUGHNUT")):
        p = deck(kind, data, colors=["#112233", "#445566", "#778899"], center_text="{total}" if kind == "donut" else None)
        exported = p.export(format="PPTX")
        (shape,) = chart_shapes(exported)
        assert shape.chart.chart_type.name == chart_type
        series = shape.chart.plots[0].series[0]
        assert list(series.values) == [15900.0, 11960.0, 2510.0]
        assert [series.points[i].format.fill.fore_color.rgb.__str__() for i in range(3)] == ["112233", "445566", "778899"]
        xml = zipfile.ZipFile(io.BytesIO(exported)).read("ppt/charts/chart1.xml").decode()
        assert '<c:showPercent val="1"/>' in xml
    texts = [s.text_frame.text for s in pptx_slide(exported).shapes if s.has_text_frame]
    assert "30,370" in texts, "the donut's center text rides over the chart as a text box"


def pptx_slide(data):
    import pptx

    return pptx.Presentation(io.BytesIO(data)).slides[0]


def test_pptx_chart_embeds_the_table_so_edit_data_works():
    p = deck("bar", {"categories": ["Jan", "Feb"], "series": [{"name": "Espresso", "values": [1200, 1100]},
                                                        {"name": "Latte", "values": [800, None]}]})
    exported = p.export(format="PPTX")
    package = zipfile.ZipFile(io.BytesIO(exported))
    names = package.namelist()
    assert "ppt/charts/chart1.xml" in names and "ppt/embeddings/Microsoft_Excel_Sheet1.xlsx" in names
    rels = package.read("ppt/charts/_rels/chart1.xml.rels").decode()
    assert "../embeddings/Microsoft_Excel_Sheet1.xlsx" in rels and "relationships/package" in rels
    assert "chart" in package.read("ppt/slides/_rels/slide1.xml.rels").decode()
    types = package.read("[Content_Types].xml").decode()
    assert "drawingml.chart+xml" in types and 'Extension="xlsx"' in types
    book = zipfile.ZipFile(io.BytesIO(package.read("ppt/embeddings/Microsoft_Excel_Sheet1.xlsx")))
    main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    strings = [si.findtext(main + "t") for si in ElementTree.fromstring(book.read("xl/sharedStrings.xml"))]
    cells = {}
    for cell in ElementTree.fromstring(book.read("xl/worksheets/sheet1.xml")).iter(main + "c"):
        raw = cell.findtext(main + "v")
        cells[cell.get("r")] = strings[int(raw)] if cell.get("t") == "s" else float(raw)
    assert cells == {"B1": "Espresso", "C1": "Latte", "A2": "Jan", "B2": 1200.0, "C2": 800.0, "A3": "Feb", "B3": 1100.0}
    chart_xml = package.read("ppt/charts/chart1.xml").decode()
    assert "Sheet1!$B$2:$B$3" in chart_xml and "Sheet1!$C$1" in chart_xml
    for name in names:
        if name.endswith((".xml", ".rels")):
            ElementTree.fromstring(package.read(name))  # every part is well-formed


def test_fixing_a_number_updates_the_native_chart_too():
    p = deck("bar", SALES)
    p.apply({"type": "chart-data", "target": "Sales", "set": [{"category": "Jun", "value": 700}]})
    (shape,) = chart_shapes(p.export(format="PPTX"))
    assert list(shape.chart.plots[0].series[0].values)[-1] == 700.0


def test_pptx_chart_follows_the_chart_settings():
    p = deck("stacked-bar", DRINKS, legend="right", gridlines=False, bar_gap=0.5, colors=["#ff0000", "#00ff00", "#0000ff"],
             value_title="Cups", category_title="Month", value_labels=True)
    exported = p.export(format="PPTX")
    xml = zipfile.ZipFile(io.BytesIO(exported)).read("ppt/charts/chart1.xml").decode()
    assert '<c:legendPos val="r"/>' in xml and "<c:majorGridlines>" not in xml
    assert '<c:gapWidth val="100"/>' in xml and '<c:overlap val="100"/>' in xml
    assert 'srgbClr val="FF0000"' in xml and '<c:dLblPos val="ctr"/>' in xml
    assert "<a:t>Cups</a:t>" in xml and "<a:t>Month</a:t>" in xml
    (shape,) = chart_shapes(exported)
    assert shape.chart.value_axis.maximum_scale == p.layer("Sales")["chart"]["summary"]["scale"]["max"]
    assert shape.chart.value_axis.tick_labels.number_format == "#,##0"


def test_rotated_chart_exports_as_shapes_with_a_note():
    p = deck("bar", SALES)
    p.apply({"type": "rotate", "target": "Sales", "value": 5})
    report = {}
    data = p.export(format="PPTX", report=report)
    assert not chart_shapes(data)
    assert report["charts"]["1"][0]["native"] is False and "shapes" in report["charts"]["1"][0]["reason"]
    slide = pptx_slide(data)
    assert any(shape.shape_type == 6 for shape in slide.shapes)  # a group of the generated shapes


def test_chart_with_effects_is_a_picture_like_any_other_group():
    p = deck("bar", SALES)
    p.apply({"type": "layer-style", "target": "Sales", "name": "drop-shadow", "settings": {"dx": 4, "dy": 4, "blur": 6}})
    report = {}
    data = p.export(format="PPTX", report=report)
    assert not chart_shapes(data) and report["raster_fallbacks"]["1"][0]["layer"] == "Sales"


def test_operation_schema_documents_chart_operations():
    variants = {v["properties"]["type"]["const"]: v for v in operation_schema()["properties"]["operations"]["items"]["oneOf"]}
    chart, data = variants["chart"], variants["chart-data"]
    assert chart["description"] and data["description"]
    assert set(chart["properties"]["kind"]["enum"]) >= {"bar", "stacked-bar", "horizontal-bar", "line", "pie", "donut", "area"}
    for key in ("categories", "series", "table", "csv", "title", "legend", "value_labels", "number_format", "colors"):
        assert chart["properties"][key]["description"]
    assert chart["properties"]["series"]["items"]["required"] == ["name", "values"]
    assert {"set", "append", "remove_categories", "add_series", "remove_series", "reload"} <= set(data["properties"])
    assert data["properties"]["set"]["items"]["properties"]["value"]["type"] == ["number", "null"]
    validate_operation({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "kind": "pie"})
    with pytest.raises(VixlError, match="kind"):
        validate_operation({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "kind": "radar"})
    with pytest.raises(VixlError, match="Unknown field"):
        validate_operation({"type": "chart", "categories": ["a"], "series": [{"name": "s", "values": [1]}], "legned": "top"})


def test_service_operations_accept_charts(tmp_path):
    import asyncio

    import jsonschema
    from vixl.interfaces import mcp_server

    async def run():
        tools = {t.name: t for t in await mcp_server(workspace=tmp_path).list_tools()}
        schema = tools["vixl_operations_apply"].inputSchema
        jsonschema.validate({"operations": [
            {"type": "chart", "name": "c", "kind": "donut", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
            {"type": "chart-data", "target": "c", "set": [{"category": "a", "value": 5}]}]}, schema)
        with pytest.raises(jsonschema.ValidationError):
            jsonschema.validate({"operations": [{"type": "chart", "kind": "radar"}]}, schema)

    asyncio.run(run())


def test_cli_syntax_compiles_to_chart_operations():
    op = compile_command(["chart", "stacked-bar", "--name", "Sales", "--csv", "data.csv", "--title", "Cups", "--colors",
                          "#f00,#0f0", "--legend", "top", "--value-labels", "true", "--no-gridlines", "--ticks", "4",
                          "--number-format", "#,##0", "--x", "40", "--width", "800"])
    assert op == {"type": "chart", "kind": "stacked-bar", "name": "Sales", "csv": "data.csv", "title": "Cups",
                  "colors": ["#f00", "#0f0"], "legend": "top", "value_labels": True, "gridlines": False, "ticks": 4,
                  "number_format": "#,##0", "x": 40.0, "width": 800}
    assert compile_command(["chart", "--target", "Sales", "--min", "0"]) == {"type": "chart", "target": "Sales", "min": 0.0}
    edit = compile_command(["chart-data", "--target", "Sales", "--set", "Dec=3330", "--set", "Jan:Latte=", "--append", "Feb=1,2",
                            "--remove-category", "Mar", "--reload"])
    assert edit == {"type": "chart-data", "target": "Sales", "set": [{"category": "Dec", "value": 3330},
                                                                      {"category": "Jan", "series": "Latte", "value": None}],
                    "append": [{"category": "Feb", "values": [1, 2]}], "remove_categories": ["Mar"], "reload": True}
    p = make("bar", {"categories": ["Nov", "Dec"], "series": [{"name": "Cups", "values": [1, 2]}]})
    p.apply(compile_command(["chart-data", "--target", "Sales", "--set", "Dec=3330"]))
    assert p.layer("Sales")["chart"]["series"][0]["values"] == [1, 3330]


def test_cli_end_to_end(tmp_path):
    import subprocess
    import sys
    import os

    env = {**os.environ, "PYTHONPATH": str(__import__("pathlib").Path(__file__).parents[1] / "src")}
    (tmp_path / "cups.csv").write_text("Month,Cups\nJan,10\nFeb,20\n")
    def run(*args):
        done = subprocess.run([sys.executable, "-m", "vixl", *args], cwd=tmp_path, capture_output=True, text=True, env=env)
        assert done.returncode == 0, done.stderr
        return json.loads(done.stdout) if done.stdout else None

    run("new", "600x400", "-o", "c.vixl")
    run("chart", "bar", "--name", "Sales", "--csv", "cups.csv", "--title", "Cups")
    (tmp_path / "cups.csv").write_text("Month,Cups\nJan,10\nFeb,40\n")
    run("chart-data", "--target", "Sales", "--reload")
    run("chart-data", "--target", "Sales", "--set", "Jan=15")
    loaded = Project.load(tmp_path / "c.vixl")
    assert loaded.layer("Sales")["chart"]["series"][0]["values"] == [15, 40]
    assert run("check", "--checks", "bounds")["passed"]
