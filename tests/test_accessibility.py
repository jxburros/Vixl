"""Alt text, language, reading order and the accessibility check (#175, scope from #561)."""

import io
import re
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image
import pytest

from vixl import Project, VixlError
from vixl.checks import check_design


def photo(p, name="photo", x=20, y=120):
    buffer = io.BytesIO()
    Image.new("RGB", (60, 40), "#3a6").save(buffer, "PNG")
    from vixl.assets import add_encoded

    asset, _ = add_encoded(p, buffer.getvalue())
    p.apply({"type": "add", "name": name, "asset": asset, "x": x, "y": y})


def poster():
    p = Project(400, 300, "white")
    p.apply([
        {"type": "text", "name": "title", "text": "Field notes", "size": 36, "x": 20, "y": 20, "color": "#111111"},
        {"type": "shape", "shape": "star", "name": "sparkle", "x": 330, "y": 20, "width": 40, "height": 40,
         "fill": "#fc0"},
    ])
    photo(p)
    return p


def findings(p, code=None, **options):
    issues = check_design(p, checks=["accessibility"], **options)["issues"]
    return [i for i in issues if i["check"] == "accessibility" and (code is None or i.get("code") == code)]


def test_alt_and_decorative_round_trip_and_validate(tmp_path):
    p = poster()
    p.apply([{"type": "alt-text", "target": "photo", "alt_text": "A fern in a clay pot"},
             {"type": "layer-intent", "target": "sparkle", "decorative": True}])
    assert p.layer("photo")["alt"] == "A fern in a clay pot" and p.layer("sparkle")["decorative"] is True
    p.save(tmp_path / "a.vixl")
    loaded = Project.load(tmp_path / "a.vixl")
    assert loaded.layer("photo")["alt"] == "A fern in a clay pot" and loaded.layer("sparkle")["decorative"]
    with pytest.raises(VixlError, match="its own text"):
        p.apply({"type": "layer-intent", "target": "title", "alt": "Title"})
    with pytest.raises(VixlError, match="not both"):
        p.apply({"type": "layer-intent", "target": "photo", "alt": "x", "decorative": True})
    p.apply({"type": "layer-intent", "target": "photo", "alt": ""})
    assert "alt" not in p.layer("photo")


def test_check_reports_language_missing_alt_small_text_and_bundles_contrast():
    p = poster()
    p.apply({"type": "text", "name": "fine", "text": "tiny print", "size": 9, "x": 20, "y": 260, "color": "#cccccc"})
    report = check_design(p, checks=["accessibility"])
    checks = {issue["check"] for issue in report["issues"]}
    assert {"accessibility", "contrast"} <= checks  # contrast is part of the bundle
    assert [i["action"] for i in findings(p, "language")] == ["fix"]
    missing = findings(p, "missing-alt")
    assert [i["layers"] for i in missing] == [["photo"]] and missing[0]["action"] == "fix"
    assert [i["layers"] for i in findings(p, "text-size")] == [["fine"]]
    assert not report["passed"]
    p.apply([{"type": "accessibility", "lang": "en-GB"},
             {"type": "layer-intent", "target": "photo", "alt": "A fern in a clay pot"},
             {"type": "layer-intent", "target": "sparkle", "role": "decoration"}])
    assert not findings(p, "language") and not findings(p, "missing-alt")
    with pytest.raises(VixlError, match="BCP 47"):
        p.apply({"type": "accessibility", "lang": "english please"})


def test_untagged_chart_image_in_a_deck_is_a_fix_finding_and_colour_only_charts_are_flagged():
    p = Project(960, 540, "white")
    p.apply([{"type": "page", "action": "add", "name": "results"}, {"type": "accessibility", "lang": "en"}])
    photo(p, "chart-image")
    assert [i["action"] for i in findings(p, "missing-alt")] == ["fix"]
    p.apply({"type": "chart", "name": "sales", "kind": "bar", "categories": ["Q1", "Q2"],
             "series": [{"name": "North", "values": [3, 5]}, {"name": "South", "values": [4, 2]}],
             "value_labels": False, "x": 300, "y": 40, "width": 400, "height": 300})
    names = [i["layers"][0] for i in findings(p, "missing-alt")]
    assert "sales" in names
    assert [i["layers"][0] for i in findings(p, "color-only-chart")] == ["sales"]
    p.apply({"type": "layer-intent", "target": "sales", "alt": "Bar chart: North grows, South falls",
             "color_vision_safe": True})
    assert "sales" not in [i["layers"][0] for i in findings(p, "missing-alt")]
    assert not findings(p, "color-only-chart")


def test_reading_order_finding_and_explicit_order():
    p = Project(400, 300, "white")
    p.apply([{"type": "accessibility", "lang": "en"},
             {"type": "text", "name": "footer", "text": "Read me last", "size": 20, "x": 20, "y": 250},
             {"type": "text", "name": "headline", "text": "Read me first", "size": 30, "x": 20, "y": 20}])
    found = findings(p, "reading-order")
    assert found and found[0]["layers"] == ["footer", "headline"] and found[0]["action"] == "review"
    p.apply({"type": "accessibility", "reading_order": ["headline", "footer"]})
    assert not findings(p, "reading-order")
    assert p.state["page_accessibility"]["reading_order"] == [p.layer("headline")["id"], p.layer("footer")["id"]]


def test_html_carries_language_title_and_alt():
    p = poster()
    p.apply([{"type": "accessibility", "lang": "fr-CA", "title": "Notes de terrain", "page_alt": "Une affiche"},
             {"type": "layer-intent", "target": "photo", "alt": "Une fougère"}])
    html = p.export(format="HTML").decode()
    assert '<html lang="fr-CA">' in html and "<title>Notes de terrain</title>" in html
    assert 'alt="Une affiche"' in html and "Une fougère" in html


def test_presenter_deck_uses_language_headings_alt_and_reading_order():
    p = Project(960, 540, "white")
    p.apply([{"type": "page", "action": "add", "name": "one"},
             {"type": "accessibility", "lang": "de", "title": "Bericht"},
             {"type": "text", "name": "title", "text": "Ergebnisse", "size": 48, "x": 40, "y": 40},
             {"type": "text", "name": "late", "text": "Zweitens", "size": 24, "x": 40, "y": 400},
             {"type": "text", "name": "early", "text": "Erstens", "size": 24, "x": 40, "y": 200}])
    photo(p, "graph", 500, 200)
    p.apply([{"type": "layer-intent", "target": "graph", "alt": "Diagramm der Umsätze"},
             {"type": "accessibility", "reading_order": ["late", "early"], "page_lang": "de-AT"},
             {"type": "page", "action": "add", "name": "two"},
             {"type": "text", "name": "t2", "text": "Danke", "size": 48, "x": 40, "y": 40}])
    html = p.export(format="HTML", presenter=True).decode()
    assert '<html lang="de"' in html and "<h1 class=\"sr\">Bericht</h1>" in html and "<h2>Ergebnisse</h2>" in html
    assert 'lang="de-AT"' in html
    readable = re.search(r'id="slide-1".*?<div class="sr" dir="auto">(.*?)</div>', html, re.S).group(1)
    assert readable.index("Zweitens") < readable.index("Erstens") and "Diagramm der Umsätze" in readable


def test_svg_gets_lang_title_desc_role_and_aria_hidden():
    p = poster()
    p.apply([{"type": "accessibility", "lang": "en-GB", "title": "Field notes", "page_alt": "A poster about ferns"},
             {"type": "layer-intent", "target": "photo", "alt": "A fern"},
             {"type": "layer-intent", "target": "sparkle", "decorative": True}])
    root = ET.fromstring(p.export(format="SVG"))
    ns = "{http://www.w3.org/2000/svg}"
    assert root.get("{http://www.w3.org/XML/1998/namespace}lang") == "en-GB" and root.get("role") == "img"
    assert root.find(f"{ns}title").text == "Field notes" and root.find(f"{ns}desc").text == "A poster about ferns"
    photo_group = root.find(".//*[@data-layer='photo']")
    assert photo_group.get("role") == "img" and photo_group.find(f"{ns}title").text == "A fern"
    assert root.find(".//*[@data-layer='sparkle']").get("aria-hidden") == "true"


def test_pptx_descr_decorative_flag_and_language():
    from pptx import Presentation

    p = poster()
    p.apply([{"type": "accessibility", "lang": "es-MX"},
             {"type": "layer-intent", "target": "photo", "alt": "Un helecho"},
             {"type": "layer-intent", "target": "sparkle", "decorative": True}])
    data = p.export(format="PPTX")
    deck = Presentation(io.BytesIO(data))
    shapes = {shape.name: shape for shape in deck.slides[0].shapes}
    assert shapes["photo"]._element.nvPicPr.cNvPr.get("descr") == "Un helecho"
    assert "decorative" in ET.tostring(shapes["sparkle"]._element).decode()
    slide = zipfile.ZipFile(io.BytesIO(data)).read("ppt/slides/slide1.xml").decode()
    assert 'lang="es-MX"' in slide and 'lang="en-US"' not in slide
    assert deck.core_properties.language == "es-MX"


def test_pdf_carries_lang_and_alt_marked_content():
    import pypdf

    p = poster()
    p.apply([{"type": "accessibility", "lang": "it", "title": "Appunti"},
             {"type": "layer-intent", "target": "photo", "alt": "Una felce"},
             {"type": "layer-intent", "target": "sparkle", "decorative": True}])
    reader = pypdf.PdfReader(io.BytesIO(p.export(format="PDF")))
    assert reader.trailer["/Root"]["/Lang"] == "it" and reader.metadata.title == "Appunti"
    content = reader.pages[0].get_contents().get_data()
    alt = "Una felce".encode("utf-16-be").hex().upper().encode()
    assert b"/Figure <</Alt <FEFF" + alt + b">>> BDC" in content and b"/Artifact BMC" in content
    assert content.count(b"EMC") >= 2


def test_cli_compiles_accessibility_and_alt():
    from vixl.commands import compile_command

    assert compile_command("accessibility --lang en-GB --title T") == {"type": "accessibility", "lang": "en-GB",
                                                                      "title": "T"}
    assert compile_command("layer-intent photo --alt 'A fern'")["alt"] == "A fern"
